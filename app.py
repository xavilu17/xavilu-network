import streamlit as st
import streamlit.components.v1 as components
import ollama
import asyncio
import edge_tts
import pypdf
from docx import Document
import openpyxl
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches
from io import BytesIO
from PIL import Image
from fpdf import FPDF
import urllib.parse
import urllib.request
import csv
import os
import base64
import pandas as pd
import plotly.express as px
from youtube_transcript_api import YouTubeTranscriptApi
import re
import requests
import tempfile
import time
import qrcode
import gc
import html
import numpy as np
import random
import sqlite3
import json
import datetime
import concurrent.futures

try: import cv2; HAS_CV2 = True
except ImportError: HAS_CV2 = False
try: import yt_dlp; HAS_YTDLP = True
except ImportError: HAS_YTDLP = False
try: from ddgs import DDGS
except ImportError:
    try: from duckduckgo_search import DDGS
    except ImportError: DDGS = None
try: from google import genai; from google.genai import types; HAS_GEMINI = True
except ImportError: HAS_GEMINI = False
try: import openai; HAS_OPENAI = True
except ImportError: HAS_OPENAI = False
try: import anthropic; HAS_ANTHROPIC = True
except ImportError: HAS_ANTHROPIC = False
try: import kaleido; HAS_KALEIDO = True
except ImportError: HAS_KALEIDO = False

# --- CONFIGURACIÓN ---
st.set_page_config(page_title="Xavilu Network Pro", page_icon="logo.png", layout="wide", initial_sidebar_state="expanded")
MAX_FILE_MB = 1000
MAX_TEXTO_IA = 60_000
GEMINI_MODEL = "gemini-3.8-flash"

if os.name == 'nt':
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

# --- CATÁLOGOS MAESTROS DE LA PLANTA ---
CATALOGO_LER = {
    "15 01 01": "Envases de papel y cartón",
    "15 01 02": "Envases de plástico",
    "15 01 04": "Envases metálicos",
    "15 01 07": "Envases de vidrio",
    "20 01 01": "Papel y cartón",
    "20 01 39": "Plásticos",
    "20 01 40": "Metales",
    "20 01 08": "Residuos biodegradables (Orgánico)",
    "17 04 05": "Hierro y acero (Chatarra)",
    "17 04 02": "Aluminio",
    "17 02 01": "Madera",
    "20 01 21*": "Tubos fluorescentes y otros residuos que contienen mercurio",
    "15 02 02*": "Absorbentes, materiales de filtración, trapos de limpieza contaminados",
    "16 06 01*": "Baterías de plomo",
    "13 02 05*": "Aceites minerales no clorados de motor",
    "OTRO": "Otro código (Especificar en notas)"
}

DESTINOS_INTERNOS = [
    "Línea 1 - Triaje Manual",
    "Línea 2 - Triaje Automático",
    "Prensa de Cartón",
    "Lavado de PET",
    "Foso de Orgánico",
    "Zona de Peligrosos",
    "Foso de Rechazo",
    "Almacén de Chatarra"
]

# --- BASE DE DATOS PLANTAS ---
def init_plant_db():
    conn = sqlite3.connect('xavilu_planta.db')
    c = conn.cursor()
    # TABLA MAESTRA: FLOTA DE VEHÍCULOS (Dashboard)
    c.execute('''CREATE TABLE IF NOT EXISTS flota_vehiculos (
        id INTEGER PRIMARY KEY AUTOINCREMENT, 
        matricula TEXT UNIQUE, 
        empresa TEXT, 
        conductor TEXT, 
        dni TEXT, 
        tara_fija REAL
    )''')
    # TABLA OPERATIVA: BÁSCULA (Visitas)
    c.execute('''CREATE TABLE IF NOT EXISTS bascula (
        id INTEGER PRIMARY KEY AUTOINCREMENT, 
        fecha TEXT, 
        matricula TEXT, 
        empresa TEXT, 
        residuo TEXT, 
        codigo_ler TEXT, 
        origen TEXT, 
        destino TEXT, 
        peso_bruto REAL, 
        tara REAL, 
        peso_neto REAL, 
        penalizacion REAL, 
        albaran TEXT, 
        doc_di TEXT, 
        inspeccion_ok INTEGER, 
        incidencia TEXT, 
        operacion TEXT
    )''')
    # TABLA TRAZABILIDAD
    c.execute('''CREATE TABLE IF NOT EXISTS lotes (
        id INTEGER PRIMARY KEY AUTOINCREMENT, 
        fecha TEXT, 
        residuo TEXT, 
        peso_in REAL, 
        estado TEXT, 
        peso_out REAL
    )''')
    # TABLA MANTENIMIENTO
    c.execute('''CREATE TABLE IF NOT EXISTS mantenimiento (
        id INTEGER PRIMARY KEY AUTOINCREMENT, 
        fecha TEXT, 
        maquina TEXT, 
        descripcion TEXT, 
        estado TEXT
    )''')
    # TABLAS SIMULADOR
    c.execute('''CREATE TABLE IF NOT EXISTS escenarios_simulacion (
        id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT, referencia TEXT, nombre TEXT, cliente TEXT, notas TEXT, parametros_json TEXT, ingresos REAL, costes REAL, beneficio REAL
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS mantenimiento_legal (
        id INTEGER PRIMARY KEY AUTOINCREMENT, planta TEXT, equipo TEXT, norma TEXT, frecuencia_dias INTEGER, ultima_ejecucion TEXT, proxima_ejecucion TEXT, responsable TEXT, notas TEXT, estado TEXT
    )''')
    c.execute('''CREATE TABLE IF NOT EXISTS inventario_movimientos (
        id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT, material TEXT, tipo TEXT, cantidad REAL, escenario_id INTEGER, notas TEXT
    )''')
    conn.commit(); conn.close()
init_plant_db()

def get_db_data(table):
    try:
        conn = sqlite3.connect('xavilu_planta.db')
        df = pd.read_sql_query(f"SELECT * FROM {table}", conn)
        conn.close(); return df
    except: return pd.DataFrame()

def insert_db(query, params):
    try:
        conn = sqlite3.connect('xavilu_planta.db')
        conn.execute(query, params)
        conn.commit(); conn.close()
    except Exception as e: st.error(f"Error DB: {e}")

# --- BASE DE DATOS HISTORIAL Y CHAT ---
def init_db():
    conn = sqlite3.connect('xavilu_memoria.db')
    conn.execute('CREATE TABLE IF NOT EXISTS historial (id INTEGER PRIMARY KEY AUTOINCREMENT, fecha TEXT, tipo TEXT, contenido TEXT)')
    conn.execute('CREATE TABLE IF NOT EXISTS chat_db (id INTEGER PRIMARY KEY AUTOINCREMENT, rol TEXT, contenido TEXT)')
    conn.commit(); conn.close()

def guardar_en_db(tipo, contenido):
    try:
        conn = sqlite3.connect('xavilu_memoria.db')
        conn.execute("INSERT INTO historial (fecha, tipo, contenido) VALUES (?, ?, ?)", (datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"), tipo, contenido))
        conn.commit(); conn.close()
    except: pass

def obtener_de_db():
    try:
        conn = sqlite3.connect('xavilu_memoria.db')
        rows = conn.execute("SELECT fecha, tipo, contenido FROM historial ORDER BY id DESC LIMIT 5").fetchall()
        conn.close(); return rows
    except: return []

def cargar_chat_db():
    try:
        conn = sqlite3.connect('xavilu_memoria.db')
        rows = conn.execute("SELECT rol, contenido FROM chat_db ORDER BY id ASC").fetchall()
        conn.close(); return [{"role": r[0], "content": r[1]} for r in rows]
    except: return []

def guardar_chat_db(rol, contenido):
    try:
        conn = sqlite3.connect('xavilu_memoria.db')
        conn.execute("INSERT INTO chat_db (rol, contenido) VALUES (?, ?)", (rol, contenido))
        conn.commit(); conn.close()
    except: pass

def limpiar_chat_db():
    try:
        conn = sqlite3.connect('xavilu_memoria.db')
        conn.execute("DELETE FROM chat_db")
        conn.commit(); conn.close()
    except: pass

init_db()

# ==========================================
# 📚 CATÁLOGOS DEL SIMULADOR DE PLANTAS
# ==========================================
CATALOGO_RESIDUOS = {
    "RSU mezcla":         {"categoria": "Doméstico",    "peligroso": False, "fracciones": ["PET", "Metales", "Cartón", "Vidrio", "Orgánico", "Otros"]},
    "Envases ligeros":    {"categoria": "Reciclable",   "peligroso": False, "fracciones": ["PET", "HDPE", "Film", "Metales", "Cartón", "Otros"]},
    "Papel-cartón":       {"categoria": "Reciclable",   "peligroso": False, "fracciones": ["Cartón", "Papel mixto", "Impropios"]},
    "Vidrio":             {"categoria": "Reciclable",   "peligroso": False, "fracciones": ["Vidrio claro", "Vidrio color", "Impropios"]},
    "Orgánico":           {"categoria": "Orgánico",     "peligroso": False, "fracciones": ["Orgánico", "Impropios"]},
    "Plásticos":          {"categoria": "Reciclable",   "peligroso": False, "fracciones": ["PET", "HDPE", "Film", "Mixto"]},
    "Metales":            {"categoria": "Reciclable",   "peligroso": False, "fracciones": ["Férricos", "Aluminio", "Cobre"]},
    "Madera":             {"categoria": "Reciclable",   "peligroso": False, "fracciones": ["Madera A", "Madera B", "Impropios"]},
    "Textil":             {"categoria": "Reciclable",   "peligroso": False, "fracciones": ["Reutilizable", "Reciclable", "Impropios"]},
    "RAEE":               {"categoria": "Especial",     "peligroso": False, "fracciones": ["Férricos", "Aluminio", "Cobre", "Plásticos", "Otros"]},
    "Voluminosos":        {"categoria": "Especial",     "peligroso": False, "fracciones": ["Madera", "Metales", "Espuma", "Otros"]},
    "Neumáticos":         {"categoria": "Especial",     "peligroso": False, "fracciones": ["Caucho", "Acero", "Textil"]},
    "Escombro":           {"categoria": "Construcción", "peligroso": False, "fracciones": ["Áridos", "Metales", "Madera", "Impropios"]},
    "Chatarra":           {"categoria": "Metálico",     "peligroso": False, "fracciones": ["Férricos", "Aluminio", "Cobre", "Otros"]},
    "Residuos peligrosos":{"categoria": "Peligroso",    "peligroso": True,  "fracciones": ["Rechazo"]},
    "Hospitalarios":      {"categoria": "Peligroso",    "peligroso": True,  "fracciones": ["Rechazo"]},
    "Aceites usados":     {"categoria": "Peligroso",    "peligroso": True,  "fracciones": ["Rechazo"]},
    "Disolventes":        {"categoria": "Peligroso",    "peligroso": True,  "fracciones": ["Rechazo"]},
}

PERFILES_PLANTA = {
    "🏭 RSU estándar":        {"residuos": ["RSU mezcla", "Envases ligeros", "Papel-cartón", "Vidrio", "Orgánico", "Metales"], "acepta_peligrosos": False, "capacidad_mes": 1500},
    "♻️ Envases ligeros":     {"residuos": ["Envases ligeros", "Plásticos"], "acepta_peligrosos": False, "capacidad_mes": 800},
    "📄 Papel-cartón":        {"residuos": ["Papel-cartón"], "acepta_peligrosos": False, "capacidad_mes": 500},
    "🍾 Vidrio":              {"residuos": ["Vidrio"], "acepta_peligrosos": False, "capacidad_mes": 600},
    "🌱 Compostaje":          {"residuos": ["Orgánico", "Madera", "Voluminosos"], "acepta_peligrosos": False, "capacidad_mes": 1000},
    "🔋 RAEE":                {"residuos": ["RAEE"], "acepta_peligrosos": False, "capacidad_mes": 300},
    "⚠️ Peligrosos":          {"residuos": ["Residuos peligrosos", "Hospitalarios", "Aceites usados", "Disolventes"], "acepta_peligrosos": True, "capacidad_mes": 200},
    "🏗️ Escombro":            {"residuos": ["Escombro"], "acepta_peligrosos": False, "capacidad_mes": 2000},
    "🌐 Multiresiduo (todo)": {"residuos": list(CATALOGO_RESIDUOS.keys()), "acepta_peligrosos": True, "capacidad_mes": 5000},
}

DURABILIDADES_DEFAULT = {
    "Orgánico": 1, "RSU mezcla": 2, "Lodos": 1, "Restos vegetales": 2,
    "RAEE (con baterías)": 7, "RAEE (sin baterías)": 30,
    "PET": 30, "HDPE": 30, "Film": 30, "Mixto": 30, "Plásticos": 30,
    "Papel mixto": 60, "Cartón": 30,
    "Vidrio claro": 180, "Vidrio color": 180, "Vidrio": 180,
    "Férricos": 180, "Aluminio": 180, "Cobre": 180, "Metales": 180, "Acero": 180,
    "Madera A": 90, "Madera B": 60, "Madera": 60,
    "Reutilizable": 60, "Reciclable": 60, "Textil": 60,
    "Caucho": 180, "Áridos": 365, "Espuma": 60, "Otros": 60, "Impropios": 5, "Rechazo": 3,
}

BIBLIOTECA_MANTENIMIENTO = [
    {"equipo": "Equipos de trabajo (general)",         "norma": "RD 1215/1997",          "frecuencia_dias": 365},
    {"equipo": "Máquinas nuevas (marcado CE)",         "norma": "RD 1644/2008",          "frecuencia_dias": 365},
    {"equipo": "Grúas / polipastos",                   "norma": "ITC MIE-AEM-2",         "frecuencia_dias": 365},
    {"equipo": "Carretillas elevadoras",               "norma": "RD 1215/1997",          "frecuencia_dias": 365},
    {"equipo": "Instalación eléctrica AT",             "norma": "RD 337/2014 (RAT)",     "frecuencia_dias": 365},
    {"equipo": "Instalación eléctrica BT",             "norma": "RD 842/2002",           "frecuencia_dias": 1460},
    {"equipo": "Aparatos a presión",                   "norma": "RD 809/2021",           "frecuencia_dias": 730},
    {"equipo": "Extintores (revisión)",                "norma": "RD 513/2017",           "frecuencia_dias": 365},
    {"equipo": "Extintores (retimbrado)",              "norma": "RD 513/2017",           "frecuencia_dias": 1825},
    {"equipo": "Sistemas contra incendios",            "norma": "RD 513/2017",           "frecuencia_dias": 90},
    {"equipo": "Detección de gases / ATEX",            "norma": "RD 681/2003",           "frecuencia_dias": 180},
    {"equipo": "Depósitos de combustible",             "norma": "MI-IP-03",              "frecuencia_dias": 1825},
    {"equipo": "Silos / almacenamiento polvo",         "norma": "ATEX zona 20/21/22",    "frecuencia_dias": 365},
    {"equipo": "Cintas transportadoras",               "norma": "Manual fabricante",     "frecuencia_dias": 180},
    {"equipo": "Trituradoras",                         "norma": "Manual fabricante",     "frecuencia_dias": 180},
    {"equipo": "Análisis de aguas",                    "norma": "AAI de la planta",      "frecuencia_dias": 90},
    {"equipo": "Control de emisiones",                 "norma": "Ley 7/2022 / AAI",      "frecuencia_dias": 180},
    {"equipo": "Revisión de básculas",                 "norma": "Reglamento metrología", "frecuencia_dias": 365},
]

# --- ESTADO DE SESIÓN ---
_defaults = {
    "password_correct": False, "rol": "ninos",
    "chat_history": cargar_chat_db(), "img_url": "", "app_code": "", "pdf_chat_history": [],
    "mermaid_code": "", "analisis_ia": "", "dl_bytes": None, "dl_filename": "", "dl_mime": "",
    "cuento_generado": "", "receta_generada": "", "menu_generado": "",
    "redes_resultado": "", "pelo_resultado": "", "deberes_resultado": "", "prl_resultado": "",
    "qr_resultado": "", "code_resultado": "", "entrevista_resultado": "",
    "presupuesto_resultado": "", "noticias_deportes": [], "resumen_deportivo": "",
    "evento_resultados": None, "office_text": "", "pdf_text": "", "clip_text": "",
    "dep_resultados": "", "partido_analisis": "", "plan_deporte": "",
    "cine_rec": "", "tv_prog": "", "entradas": "", "trivia": "", "yt_text": "", "web_text": "",
    "letra_cancion": "", "titulo_cancion": "", "portada_cancion": "",
    "doc_opcion": "Resumen Ejecutivo",
    "global_context": "", "global_file_name": "",
    "agente_activo": None, "agente_chat_history": [],
    "debate_resultado": "", "img_prompt_en": "",
    "podcast_script": "", "podcast_ideas": [],
    "sidebar_visible": True,
    "whatsapp_number": "",
    "office_borrador": "", "pdf_borrador": "",
    "plant_borrador": "", "work_result": "",
    "pod_tema_input": "", "pod_sugs": [],
    "api_key_gemini_manual": "", "api_key_groq_manual": "", "api_key_silicon_manual": "", 
    "api_key_mistral_manual": "", "api_key_openrouter_manual": "",
    "galeria_imgs": [], "office_estilo": "Normal",
    "sim_config_planta": {"nombre": "Planta Demo Xavilu", "plantilla": "🏭 RSU estándar", "capacidad_mes": 1500, "acepta_peligrosos": False, "residuos_habilitados": ["RSU mezcla", "Envases ligeros", "Papel-cartón", "Vidrio", "Orgánico", "Metales"]},
    "sim_tipos_residuo": [],
    "sim_fracciones": [],
    "sim_matriz": {},
    "sim_stock_inicial": {},
    "sim_durabilidades": dict(DURABILIDADES_DEFAULT),
    "sim_costes": {"tarifa": 35.0, "variable": 18.0, "fijos": 12000.0, "rechazo": 65.0},
    "sim_resultado_actual": None,
    "sim_params_actuales": None,
    "sim_avisos": [],
    "sim_modo_recepcion": "Camiones detallados",
    "sim_camiones": [],
}
for _k, _v in _defaults.items():
    if _k not in st.session_state: st.session_state[_k] = _v

# ==========================================
# 🔐 SISTEMA DE LOGIN (ADMIN Y NIÑOS)
# ==========================================
def check_password():
    if st.session_state.get("password_correct", False):
        return True
    st.markdown("## 🔐 Xavilu Network")
    st.caption("Introduce la contraseña para acceder")
    with st.form("login_form"):
        pwd = st.text_input("Contraseña", type="password")
        submitted = st.form_submit_button("Entrar", use_container_width=True)
    if submitted:
        try:
            pwd_admin = st.secrets.get("PASSWORD", "admin")
            pwd_ninos = st.secrets.get("PASSWORD_NINOS", "ninos")
        except:
            pwd_admin = "admin"
            pwd_ninos = "ninos"
        if pwd == pwd_admin:
            st.session_state["password_correct"] = True
            st.session_state["rol"] = "admin"
            st.rerun()
        elif pwd == pwd_ninos:
            st.session_state["password_correct"] = True
            st.session_state["rol"] = "ninos"
            st.rerun()
        else:
            st.error("❌ Contraseña incorrecta. Inténtalo de nuevo.")
    return False

if not check_password():
    st.stop()

# ==========================================
# 🎨 CSS NEGRO ESPACIAL Y ARREGLOS MÓVIL
# ==========================================
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Orbitron:wght@400;700&family=Rajdhani:wght@500;700&display=swap');
.stApp { background-color: #000000 !important; background-image: radial-gradient(circle at 50% 50%, #111111 0%, #000000 100%) !important; color: #f7dfa3 !important; font-family: 'Rajdhani', sans-serif !important; }
header[data-testid="stHeader"] { visibility: hidden !important; }
#MainMenu { visibility: hidden !important; }
footer { visibility: hidden !important; }
[data-testid="stAppDeployButton"] { display: none !important; }
[data-testid="stDecoration"] { display: none !important; }
[data-testid="stToolbar"] { visibility: hidden !important; }
[data-testid="stStatusWidget"] { visibility: hidden !important; }

h1 { font-family: 'Orbitron', sans-serif !important; color: #FFD700 !important; text-shadow: 0 0 15px #FF8C00, 0 0 25px #FFD700, 0 0 35px #FF8C00 !important; text-align: center; letter-spacing: 3px; }
h2, h3 { font-family: 'Orbitron', sans-serif !important; color: #ffffff !important; text-shadow: 0 0 10px #FF007F, 0 0 20px #00f0ff, 0 0 30px #FF007F !important; text-align: center; letter-spacing: 3px; }
h4 { font-family: 'Orbitron', sans-serif !important; color: #00f0ff !important; text-shadow: 0 0 8px #00f0ff, 0 0 15px #FF007F !important; letter-spacing: 1px; }
[data-testid="stWidgetLabel"] p { color: #00f0ff !important; font-family: 'Orbitron', sans-serif !important; font-size: 13px !important; font-weight: bold !important; text-transform: uppercase !important; background: rgba(0, 240, 255, 0.1) !important; padding: 5px 10px !important; border-left: 4px solid #FF007F !important; border-radius: 4px !important; display: inline-block !important; margin-bottom: 8px !important; }
[data-testid="stSidebar"] { background-color: rgba(5, 5, 10, 0.98) !important; border-right: 2px solid #00f0ff !important; box-shadow: 5px 0 15px rgba(0, 240, 255, 0.2) !important; }
[data-testid="stRadio"] > div[role="radiogroup"] label div p, [data-testid="stRadio"] > div[role="radiogroup"] label div span { color: #ffffff !important; font-size: 15px !important; font-weight: 600 !important; }
[data-testid="stFileUploader"] { background-color: #0a0a0a !important; border: 1px solid #00f0ff !important; border-radius: 8px !important; padding: 10px !important; }
[data-testid="stFileUploaderDropzone"] { background-color: #0a0a0a !important; color: #ffffff !important; }
[data-testid="stFileUploaderDropzone"] p { color: #00f0ff !important; font-weight: bold !important; }
[data-testid="stExpander"] { background-color: #0a0a0a !important; border: 1px solid #00f0ff !important; border-radius: 8px !important; }
[data-testid="stExpander"] summary { color: #00f0ff !important; font-weight: bold !important; }
div.stButton > button[kind="secondary"] { background-color: #0a0a0a !important; color: #ffffff !important; font-family: 'Orbitron', sans-serif !important; font-size: 13px !important; font-weight: bold !important; border: 1px solid #00f0ff !important; border-radius: 25px !important; padding: 10px 15px !important; margin-bottom: 5px !important; transition: all 0.3s ease !important; box-shadow: 0 0 10px rgba(0, 240, 255, 0.2) !important; width: 100% !important; }
div.stButton > button[kind="secondary"]:hover { background-color: rgba(0, 240, 255, 0.15) !important; color: #00f0ff !important; box-shadow: 0 0 20px rgba(0, 240, 255, 0.7) !important; transform: translateY(-2px) !important; }
div.stButton > button[kind="primary"] { background: linear-gradient(135deg, #00f0ff 0%, #FF007F 100%) !important; color: #000000 !important; font-family: 'Orbitron', sans-serif !important; font-size: 13px !important; font-weight: 900 !important; border: 2px solid #ffffff !important; border-radius: 25px !important; padding: 10px 15px !important; margin-bottom: 5px !important; transition: all 0.3s ease !important; box-shadow: 0 0 18px rgba(0, 240, 255, 0.9), 0 0 32px rgba(255, 0, 127, 0.7) !important; transform: translateY(-1px) !important; width: 100% !important; }
.stTextArea textarea, .stTextInput input, .stSelectbox [data-baseweb="select"] { background-color: #0a0a0a !important; border: 1px solid #00f0ff !important; border-radius: 8px !important; color: #ffffff !important; font-size: 16px !important; box-shadow: inset 0 0 10px rgba(0, 240, 255, 0.1) !important; }
.stTextArea textarea:focus, .stTextInput input:focus { border: 1px solid #FF007F !important; box-shadow: 0 0 15px rgba(255, 0, 127, 0.5) !important; }
[data-testid="stTextInput"] div, [data-testid="stTextInput"] input { background-color: #0a0a0a !important; color: #ffffff !important; }
[data-testid="stTextInput"] > div > div { border: none !important; box-shadow: none !important; }
[data-testid="stTextInput"] input { caret-color: #00f0ff !important; border: 1px solid #00f0ff !important; border-radius: 10px !important; padding: 10px 14px !important; font-size: 15px !important; outline: none !important; box-shadow: 0 0 10px rgba(0, 240, 255, 0.2) !important; transition: all 0.3s ease !important; width: 100% !important; }
[data-testid="stTextInput"] input:focus { border: 1px solid #FF007F !important; box-shadow: 0 0 18px rgba(255, 0, 127, 0.6) !important; }
[data-testid="stTextInput"] input::placeholder { color: #888888 !important; opacity: 1 !important; }
div.stButton > button:not([kind="primary"]):not([kind="secondary"]), div[data-testid="stDownloadButton"] > button { background: #0a0a0a !important; color: #00f0ff !important; border: 1px solid #00f0ff !important; border-radius: 25px !important; font-family: 'Orbitron', sans-serif !important; font-weight: bold !important; padding: 10px 20px !important; transition: all 0.3s !important; box-shadow: 0 0 10px rgba(0, 240, 255, 0.3) !important; }
div.stButton > button:not([kind="primary"]):not([kind="secondary"]):hover, div[data-testid="stDownloadButton"] > button:hover { background: linear-gradient(135deg, #FF007F 0%, #00f0ff 100%) !important; color: #000 !important; border: 1px solid #fff !important; transform: translateY(-2px); box-shadow: 0 0 20px rgba(255, 0, 127, 0.8) !important; }
[data-testid="stChatMessage"] { background: rgba(10, 10, 15, 0.8) !important; border: 1px solid #FF007F !important; border-radius: 15px !important; padding: 10px; box-shadow: 0 0 15px rgba(255, 0, 127, 0.3) !important; }
.stAlert { border-radius: 10px; font-weight: bold; background-color: rgba(0, 240, 255, 0.1) !important; border-left: 5px solid #00f0ff !important; }
[data-testid="stVerticalBlock"] > div { padding-bottom: 0.5rem !important; }

@media (max-width: 768px) {
    .block-container { padding-top: 1rem !important; padding-left: 1rem !important; padding-right: 1rem !important; }
    div.stButton > button { padding: 8px 10px !important; font-size: 12px !important; }
}
</style>
""", unsafe_allow_html=True)

def inyectar_js_fixer():
    components.html("<script>(function(){var doc=window.parent.document;function f(){doc.querySelectorAll('[data-testid=\"stTextInput\"] *').forEach(function(el){if(el.tagName.toLowerCase()==='input'||el.tagName.toLowerCase()==='div'){el.style.setProperty('background-color','#0a0a0a','important');el.style.setProperty('color','#ffffff','important');}});}f();setInterval(f,300);})();</script>", height=0)
inyectar_js_fixer()

# --- FUNCIONES BASE ---
def buscar_en_internet(query):
    if DDGS is None: return []
    try:
        res = DDGS().text(query, region='es-es', max_results=5)
        return [{"titulo": r.get("title", ""), "url": r.get("href", ""), "resumen": r.get("body", "")} for r in res]
    except: return []

def buscar_noticias_deportivas(deporte, liga):
    if DDGS is None: return []
    try:
        res = DDGS().news(f"últimas noticias {deporte} {liga} hoy", max_results=5)
        return [{"titulo": r.get("title", ""), "url": r.get("url", ""), "resumen": r.get("body", ""), "fecha": r.get("date", "")} for r in res]
    except: return []

def leer_archivo_universal(upload, modelo_ia):
    if upload.size > MAX_FILE_MB * 1024 * 1024: return ""
    nombre = upload.name.lower()
    if nombre.endswith(".txt"): return upload.read().decode("utf-8", errors="ignore")
    if nombre.endswith(".pdf"): return "\n".join([p.extract_text() for p in pypdf.PdfReader(BytesIO(upload.getvalue())).pages if p.extract_text()])
    if nombre.endswith(".docx"): return "\n".join([p.text for p in Document(upload).paragraphs if p.text])
    return "Formato no reconocido."

def cajita_adjuntar_archivo(key_suffix, modelo, api_key):
    with st.expander("📎 Adjuntar archivo para esta acción (Opcional)"):
        archivo = st.file_uploader("Sube PDF, Word, TXT o Imagen", type=["pdf", "txt", "docx", "png", "jpg", "jpeg"], key=f"adj_{key_suffix}")
        if archivo:
            nombre = archivo.name.lower()
            if nombre.endswith((".png", ".jpg", ".jpeg")):
                with st.spinner("Analizando imagen..."):
                    ctx = analizar_archivo_con_gemini(archivo.getvalue(), "image/jpeg", "Describe esta imagen o extrae el texto (OCR).", api_key)
                    st.success("✅ Imagen leída."); return f"\n[TEXTO DE IMAGEN]:\n{ctx}\n"
            else:
                texto = leer_archivo_universal(archivo, modelo)
                st.success(f"✅ Archivo '{archivo.name}' leído."); return f"\n[CONTENIDO DE '{archivo.name}']:\n{texto[:10000]}\n"
    return ""

@st.cache_resource
def get_gemini_client(api_key): return genai.Client(api_key=api_key)

def analizar_archivo_con_gemini(archivo_bytes, mime_type, prompt, api_key_gemini):
    if not HAS_GEMINI: return "[ERROR] Falta google-genai."
    for i in range(3):
        try:
            client = get_gemini_client(api_key_gemini)
            parte = types.Part.from_bytes(data=archivo_bytes, mime_type=mime_type)
            return client.models.generate_content(model=GEMINI_MODEL, contents=[prompt, parte]).text
        except Exception as e:
            if "503" in str(e) and i < 2: time.sleep(3); continue
            else: return f"[ERROR GEMINI] {e}"

def llamar_ia(modelo, prompt, api_key_gemini="", historial_chat=None, api_key_openai="", api_key_anthropic=""):
    prompt_final = prompt[:MAX_TEXTO_IA]
    if st.session_state.get("global_context"):
        prompt_final = f"[ARCHIVO ADJUNTO: {st.session_state.global_file_name}]\nCONTENIDO:\n{st.session_state.global_context}\n\n[INSTRUCCIÓN]:\n{prompt_final}\n\n(Responde teniendo en cuenta el archivo)."
    if historial_chat:
        contexto_chat = "\n".join([f"{msg['role'].upper()}: {msg['content']}" for msg in historial_chat[-10:]])
        prompt_final = f"Historial:\n{contexto_chat}\n\nNuevo mensaje:\n{prompt_final}\n\nResponde."

    try: key_gemini = st.session_state.api_key_gemini_manual or st.secrets.get("GEMINI_API_KEY", api_key_gemini)
    except: key_gemini = st.session_state.api_key_gemini_manual or api_key_gemini
    try: key_groq = st.session_state.api_key_groq_manual or st.secrets.get("GROQ_API_KEY", "")
    except: key_groq = st.session_state.api_key_groq_manual or ""
    try: key_silicon = st.session_state.api_key_silicon_manual or st.secrets.get("SILICONFLOW_API_KEY", "")
    except: key_silicon = st.session_state.api_key_silicon_manual or ""
    try: key_mistral = st.session_state.api_key_mistral_manual or st.secrets.get("MISTRAL_API_KEY", "")
    except: key_mistral = st.session_state.api_key_mistral_manual or ""
    try: key_openrouter = st.session_state.api_key_openrouter_manual or st.secrets.get("OPENROUTER_API_KEY", "")
    except: key_openrouter = st.session_state.api_key_openrouter_manual or ""

    res_exitosa = False
    res = None

    if key_gemini.strip() and HAS_GEMINI and not res_exitosa:
        for intento in range(2):
            try:
                res = get_gemini_client(key_gemini).models.generate_content(model=GEMINI_MODEL, contents=prompt_final).text
                res_exitosa = True
                break
            except Exception as e:
                if "429" in str(e) or "RESOURCE_EXHAUSTED" in str(e): st.warning("🔋 Cuota Gemini agotada. Probando Groq..."); break 
                elif "503" in str(e) and intento < 1: time.sleep(2); continue
                else: st.warning(f"⚠️ Error Gemini: {str(e)[:50]}. Probando Groq..."); break

    if key_groq.strip() and HAS_OPENAI and not res_exitosa:
        for m_groq in ["llama-3.3-70b-versatile", "llama-3.1-8b-instant"]:
            try:
                client = openai.OpenAI(api_key=key_groq, base_url="https://api.groq.com/openai/v1")
                res = client.chat.completions.create(model=m_groq, messages=[{"role": "user", "content": prompt_final}]).choices[0].message.content
                res_exitosa = True
                break
            except Exception as e: 
                if "404" in str(e): continue
                st.warning(f"⚠️ Error Groq ({m_groq}): {str(e)[:50]}. Probando SiliconFlow..."); break

    if key_silicon.strip() and HAS_OPENAI and not res_exitosa:
        try:
            client = openai.OpenAI(api_key=key_silicon.strip(), base_url="https://api.siliconflow.cn/v1")
            res = client.chat.completions.create(model="Qwen/Qwen2.5-7B-Instruct", messages=[{"role": "user", "content": prompt_final}]).choices[0].message.content
            res_exitosa = True
        except Exception as e: st.warning(f"⚠️ Error SiliconFlow: {str(e)[:50]}. Probando Mistral...")

    if key_mistral.strip() and HAS_OPENAI and not res_exitosa:
        try:
            client = openai.OpenAI(api_key=key_mistral, base_url="https://api.mistral.ai/v1")
            res = client.chat.completions.create(model="mistral-small-latest", messages=[{"role": "user", "content": prompt_final}]).choices[0].message.content
            res_exitosa = True
        except Exception as e: st.warning(f"⚠️ Error Mistral: {str(e)[:50]}. Probando OpenRouter...")

    if key_openrouter.strip() and HAS_OPENAI and not res_exitosa:
        for m_or in ["meta-llama/llama-3.1-8b-instruct:free", "google/gemma-2-9b-it:free"]:
            try:
                client = openai.OpenAI(api_key=key_openrouter, base_url="https://openrouter.ai/api/v1")
                res = client.chat.completions.create(model=m_or, messages=[{"role": "user", "content": prompt_final}]).choices[0].message.content
                res_exitosa = True
                break
            except Exception as e: 
                if "404" in str(e): continue
                st.warning(f"⚠️ Error OpenRouter ({m_or}): {str(e)[:50]}. Probando Ollama local..."); break

    if res_exitosa and res:
        return res

    try:
        client = ollama.Client(host="http://127.0.0.1:11434", timeout=600)
        return client.chat(model=modelo, messages=[{"role": "user", "content": prompt_final}], options={"temperature": 0.2})["message"]["content"].split("</think>")[-1].strip()
    except Exception as e: return f"[ERROR CRÍTICO] Todas las IAs fallaron o se agotaron. Detalle Ollama: {e}"

def obtener_ruta_fuente():
    for ruta in ["C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/Library/Fonts/Arial.ttf"]:
        if os.path.exists(ruta): return ruta
    return None

def generar_pdf_pro(titulo, contenido):
    pdf = FPDF(); pdf.add_page(); font_path = obtener_ruta_fuente()
    try:
        url_img = f"https://image.pollinations.ai/prompt/{urllib.parse.quote('Book cover for ' + titulo + ', cyberpunk')}?width=512&height=512&nologo=true"
        req = urllib.request.Request(url_img, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req) as response: img_data = response.read()
        with open("temp_cover.jpg", "wb") as f: f.write(img_data)
        pdf.image("temp_cover.jpg", x=40, y=50, w=130); pdf.ln(150); os.remove("temp_cover.jpg")
    except: pass
    if font_path: pdf.add_font("Arial", "", font_path, uni=True); pdf.set_font("Arial", size=11)
    else: pdf.set_font("Helvetica", size=11); contenido = contenido.encode('latin-1', 'replace').decode('latin-1')
    pdf.set_font("Helvetica", "B", 24); pdf.multi_cell(0, 10, txt=titulo, align='C')
    pdf.set_font("Helvetica", "I", 12); pdf.cell(0, 10, txt="Generado por Xavilu", align='C', ln=True)
    pdf.set_draw_color(0, 240, 255); pdf.line(10, pdf.get_y() + 5, 200, pdf.get_y() + 5); pdf.ln(15)
    if font_path: pdf.set_font("Arial", size=11)
    else: pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 8, txt=contenido); return bytes(pdf.output())

def generar_word(texto):
    doc = Document(); doc.add_heading("Documento Xavilu", 0)
    for p in texto.split("\n"):
        if p.strip(): doc.add_paragraph(p.strip())
    out = BytesIO(); doc.save(out); out.seek(0); return out.getvalue()

def generar_excel(texto_csv):
    wb = openpyxl.Workbook(); ws = wb.active
    for r_i, row in enumerate(csv.reader(texto_csv.replace("```csv", "").replace("```", "").strip().split('\n')), 1):
        for c_i, val in enumerate(row, 1):
            cell = ws.cell(row=r_i, column=c_i, value=val)
            if r_i == 1: cell.font = Font(bold=True)
    out = BytesIO(); wb.save(out); out.seek(0); return out.getvalue()

def generar_pptx(titulo, texto):
    prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text = titulo
    for bloque in [b.strip() for b in texto.split("\n\n") if b.strip()]:
        slide = prs.slides.add_slide(prs.slide_layouts[1]); lineas = bloque.split("\n")
        if slide.shapes.title: slide.shapes.title.text = lineas[0][:50]
        contenido_texto = "\n".join(lineas[1:]) if len(lineas) > 1 else ""
        tiene_ph = False
        for shape in slide.placeholders:
            if shape.placeholder_format.idx == 1: shape.text = contenido_texto; tiene_ph = True
        if not tiene_ph:
            txBox = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(5)); txBox.text_frame.text = contenido_texto
    out = BytesIO(); prs.save(out); out.seek(0); return out.getvalue()

def botones_descarga(texto, nombre_por_defecto="documento"):
    if texto:
        st.markdown("---")
        key_u = f"{nombre_por_defecto}".replace(" ", "_")
        nombre_elegido = st.text_input("📝 Nombre del archivo (sin extensión):", value=nombre_por_defecto, key=f"name_{key_u}")
        nombre_limpio = re.sub(r'[^a-zA-Z0-9_\- ]', '', nombre_elegido).strip().replace(' ', '_') or nombre_por_defecto
        
        c1, c2, c3, c4, c5 = st.columns(5)
        with c1: st.download_button("📥 TXT", texto.encode('utf-8'), f"{nombre_limpio}.txt", "text/plain", key=f"tx_{key_u}", use_container_width=True)
        with c2: st.download_button("📄 PDF", generar_pdf_pro(nombre_limpio, texto), f"{nombre_limpio}.pdf", "application/pdf", key=f"pd_{key_u}", use_container_width=True)
        with c3: st.download_button("📘 Word", generar_word(texto), f"{nombre_limpio}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", key=f"wd_{key_u}", use_container_width=True)
        with c4: st.download_button("📊 Excel", generar_excel(texto), f"{nombre_limpio}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=f"ex_{key_u}", use_container_width=True)
        with c5:
            if st.button("📲 WhatsApp", key=f"wa_{key_u}", use_container_width=True):
                num = st.session_state.get("whatsapp_number", "").strip()
                if num:
                    url_wa = f"https://wa.me/{num}?text={urllib.parse.quote(texto[:1500])}"
                    components.html(f"<script>window.open('{url_wa}', '_blank');</script>", height=0)
                    st.success("✅ Abriendo WhatsApp...")
                else:
                    st.error("⚠️ Configura tu número en la barra lateral para usar WhatsApp.")

def run_async_safe(coro):
    def _run(): return asyncio.run(coro)
    with concurrent.futures.ThreadPoolExecutor() as pool: return pool.submit(_run).result()

def generar_audio_sincrono(texto, voz):
    async def _gen():
        buf = BytesIO()
        async for chunk in edge_tts.Communicate(texto, voz).stream():
            if chunk["type"] == "audio": buf.write(chunk["data"])
        return buf.getvalue()
    try: return run_async_safe(_gen())
    except Exception as e: st.error(f"Error TTS: {e}"); return None

def generar_podcast_3_voces(guion, v1, v2, v3):
    async def _gen():
        buf = BytesIO()
        texto_procesado = re.sub(r'\(risas\)|\[risas\]|\*risas\*', 'jajaja', guion, flags=re.IGNORECASE)
        texto_procesado = re.sub(r'\(pausa\)|\[pausa\]', '...', texto_procesado, flags=re.IGNORECASE)
        texto_procesado = re.sub(r'[*#_]', '', texto_procesado)
        
        lineas = texto_procesado.split('\n')
        voz_actual = v1
        for linea in lineas:
            linea = linea.strip()
            if not linea: continue
            match = re.search(r'^Locutor\s*(\d)', linea, re.IGNORECASE)
            if match:
                num = int(match.group(1))
                if num == 1: voz_actual = v1
                elif num == 2: voz_actual = v2
                elif num == 3: voz_actual = v3
                texto_a_leer = re.sub(r'^Locutor\s*\d\s*:?', '', linea, flags=re.IGNORECASE).strip()
            else:
                texto_a_leer = linea
            if not texto_a_leer or texto_a_leer.lower() in ["elvira", "jorge", "tomas", "carlos"]: continue
            try:
                texto_final = texto_a_leer + ", "
                async for chunk in edge_tts.Communicate(texto_final, voz_actual).stream():
                    if chunk["type"] == "audio": buf.write(chunk["data"])
            except Exception: pass
        return buf.getvalue()
    try: return run_async_safe(_gen())
    except Exception as e: st.error(f"Error crítico en generación 3 voces: {e}"); return None

def campo_con_ideas(label, key_prefix, contexto_ia, modelo, api_key, placeholder="...", api_key_openai="", api_key_anthropic=""):
    input_key = f"{key_prefix}_input"
    if input_key not in st.session_state: st.session_state[input_key] = ""
    texto_ingresado = st.text_area(label, key=input_key, placeholder=placeholder, height=100)
    
    if st.button("💡 Generar 3 Ideas con IA (basadas en mi texto)", key=f"{key_prefix}_btn_ideas", use_container_width=True):
        with st.spinner("Pensando ideas..."):
            prompt = f"Basándote en esta idea: '{texto_ingresado}', dame 3 enfoques creativos o mejoras. Solo lista separada por saltos de línea, sin numeración."
            res = llamar_ia(modelo, prompt, api_key, api_key_openai=api_key_openai, api_key_anthropic=api_key_anthropic)
            st.session_state[f"{key_prefix}_sugs"] = [l.strip("-*1234567890. ") for l in res.split('\n') if len(l.strip()) > 5][:3]

    if st.session_state.get(f"{key_prefix}_sugs"):
        st.markdown("✨ **Selecciona las ideas que quieras añadir a tu texto (es opcional):**")
        seleccionadas = []
        for i, idea in enumerate(st.session_state[f"{key_prefix}_sugs"]):
            if st.checkbox(idea, key=f"{key_prefix}_chk_{i}"):
                seleccionadas.append(idea)

        c_add, c_discard = st.columns(2)
        with c_add:
            if st.button("➕ Añadir ideas seleccionadas al texto", key=f"{key_prefix}_add_selected", use_container_width=True, type="primary"):
                if seleccionadas:
                    texto_actual = st.session_state[input_key].strip()
                    nuevo_texto = "\n\n".join(seleccionadas)
                    if texto_actual:
                        st.session_state[input_key] = f"{texto_actual}\n\n{nuevo_texto}"
                    else:
                        st.session_state[input_key] = nuevo_texto
                for j in range(len(st.session_state[f"{key_prefix}_sugs"])):
                    k = f"{key_prefix}_chk_{j}"
                    if k in st.session_state: del st.session_state[k]
                del st.session_state[f"{key_prefix}_sugs"]
                st.rerun()
        with c_discard:
            if st.button("🗑️ Descartar ideas", key=f"{key_prefix}_discard", use_container_width=True):
                for j in range(len(st.session_state[f"{key_prefix}_sugs"])):
                    k = f"{key_prefix}_chk_{j}"
                    if k in st.session_state: del st.session_state[k]
                del st.session_state[f"{key_prefix}_sugs"]
                st.rerun()
                
    return texto_ingresado

def boton_superpoder_borrador(texto_base, instruccion, key_suffix, modelo, api_key, state_key, api_key_openai="", api_key_anthropic=""):
    if st.button(f"🚀 Superpoder: {instruccion}", key=f"super_{key_suffix}", type="secondary", use_container_width=True):
        if not texto_base: return
        with st.spinner("Aplicando superpoder..."):
            prompt = f"Actúa como experto. {instruccion}. Devuelve SOLO el texto mejorado, sin introducciones:\n\n{texto_base}"
            res = llamar_ia(modelo, prompt, api_key, api_key_openai=api_key_openai, api_key_anthropic=api_key_anthropic)
            if "[ERROR" in res: st.error(res)
            else: st.session_state[state_key] = res; st.rerun()

def chat_input_pro(key, placeholder="Mensaje..."):
    t_k = f"{key}_txt"; m_k = f"{key}_msg"
    if t_k not in st.session_state: st.session_state[t_k] = ""
    def _send():
        msg = st.session_state.get(t_k, "").strip()
        if msg: st.session_state[m_k] = msg; st.session_state[t_k] = ""
    c1, c2 = st.columns([6, 1])
    with c1: st.text_input("msg", placeholder=placeholder, key=t_k, label_visibility="collapsed")
    with c2: st.button("➤", key=f"{key}_btn", on_click=_send, use_container_width=True, type="primary")
    return st.session_state.pop(m_k, None)

def custom_tabs(tab_names, key_name, default_idx=0):
    s_key = f"tab_{key_name}"
    if s_key not in st.session_state: st.session_state[s_key] = default_idx
    for i in range(0, len(tab_names), 5):
        chunk = tab_names[i:i+5]; cols = st.columns(5) 
        for j, name in enumerate(chunk):
            idx = i + j
            with cols[j]:
                if st.button(name, key=f"btn_{key_name}_{idx}", use_container_width=True, type="primary" if st.session_state[s_key] == idx else "secondary"):
                    st.session_state[s_key] = idx; st.rerun()
    return st.session_state[s_key]

# ==========================================
# 🧠 MOTOR DE SIMULACIÓN (NUEVAS FUNCIONES)
# ==========================================
def evaluar_durabilidad(stock_dict, durabilidades):
    alertas = []
    for mat, cant in stock_dict.items():
        if cant <= 0: continue
        dur = durabilidades.get(mat, 60)
        dias_estimados = int(cant / 5) if cant > 10 else 1 
        if dias_estimados >= dur:
            alertas.append({"material": mat, "edad": dias_estimados, "durabilidad": dur, "estado": "Vencido"})
        elif dias_estimados >= dur * 0.7:
            alertas.append({"material": mat, "edad": dias_estimados, "durabilidad": dur, "estado": "Próximo"})
    return alertas

def calcular_simulacion(params):
    tipos_residuo = params.get("tipos_residuo", [])
    fracciones_cat = params.get("fracciones", [])
    matriz = params.get("matriz", {})
    costes = params.get("costes", {})
    stock_inicial = params.get("stock_inicial", {})
    
    total_ton = sum(t.get("Toneladas", 0) for t in tipos_residuo)
    if total_ton == 0: return None
    
    res_por_residuo = []
    produccion = {f.get("Fracción", ""): 0 for f in fracciones_cat}
    rechazo_total = 0
    perdidas_total = 0
    
    for t in tipos_residuo:
        ton = t.get("Toneladas", 0)
        if ton <= 0: continue
        rec = ton * (t.get("Recuperación (%)", 0) / 100)
        rech = ton * (t.get("Rechazo (%)", 0) / 100)
        perd = ton * (t.get("Pérdidas (%)", 0) / 100)
        
        rechazo_total += rech
        perdidas_total += perd
        
        comp_res = matriz.get(t.get("Residuo", ""), {})
        for f_cat in fracciones_cat:
            f_nom = f_cat.get("Fracción", "")
            pct_f = comp_res.get(f_nom, 0)
            cant_f = rec * (pct_f / 100)
            produccion[f_nom] += cant_f
            
        res_por_residuo.append({"residuo": t.get("Residuo", ""), "recuperado": rec, "rechazo": rech, "perdidas": perd})
        
    stock_final = {}
    ingresos_ventas = 0
    expedido_total = {}
    
    for f_cat in fracciones_cat:
        f_nom = f_cat.get("Fracción", "")
        p_in = produccion.get(f_nom, 0)
        exp_pct = f_cat.get("% Expedido", 0) / 100
        cant_exp = p_in * exp_pct
        expedido_total[f_nom] = cant_exp
        s_ini = stock_inicial.get(f_nom, 0)
        stock_final[f_nom] = s_ini + p_in - cant_exp
        ingresos_ventas += cant_exp * f_cat.get("Precio Venta (€/t)", 0)

    ing_tratamiento = total_ton * costes.get("tarifa", 0)
    coste_proceso = total_ton * costes.get("variable", 0)
    coste_rechazo = rechazo_total * costes.get("rechazo", 0)
    costes_totales = costes.get("fijos", 0) + coste_proceso + coste_rechazo
    ingresos_totales = ing_tratamiento + ingresos_ventas
    beneficio = ingresos_totales - costes_totales
    
    avisos_dur = evaluar_durabilidad(stock_final, params.get("durabilidades", {}))
    
    return {
        "total_ton": total_ton,
        "rechazo_total": rechazo_total,
        "perdidas_total": perdidas_total,
        "produccion": produccion,
        "expedido": expedido_total,
        "stock_final": stock_final,
        "ingresos_ventas": ingresos_ventas,
        "ing_tratamiento": ing_tratamiento,
        "coste_proceso": coste_proceso,
        "coste_rechazo": coste_rechazo,
        "costes_totales": costes_totales,
        "ingresos_totales": ingresos_totales,
        "beneficio": beneficio,
        "avisos_durabilidad": avisos_dur,
        "res_por_residuo": res_por_residuo
    }

# ==========================================
# 🏭 FUNCIONES UI SIMULADOR
# ==========================================
def render_simular():
    st.markdown("#### 🧮 Simulador Económico-Operativo")
    
    # 0. PERFIL DE PLANTA
    with st.expander("⚙️ 0. Perfil de la Planta", expanded=False):
        c1, c2, c3 = st.columns(3)
        with c1:
            plantilla = st.selectbox("Cargar plantilla de planta:", list(PERFILES_PLANTA.keys()), key="sim_plantilla_select")
            if st.button("🎲 Aplicar plantilla", use_container_width=True):
                p = PERFILES_PLANTA[plantilla]
                st.session_state.sim_config_planta["nombre"] = f"Planta {plantilla.split(' ')[1] if ' ' in plantilla else plantilla}"
                st.session_state.sim_config_planta["capacidad_mes"] = p["capacidad_mes"]
                st.session_state.sim_config_planta["acepta_peligrosos"] = p["acepta_peligrosos"]
                st.session_state.sim_config_planta["residuos_habilitados"] = p["residuos"].copy()
                st.rerun()
        with c2:
            st.session_state.sim_config_planta["nombre"] = st.text_input("Nombre de la planta:", value=st.session_state.sim_config_planta["nombre"])
            st.session_state.sim_config_planta["capacidad_mes"] = st.number_input("Capacidad autorizada (t/mes):", min_value=10, value=st.session_state.sim_config_planta["capacidad_mes"], step=100)
        with c3:
            st.session_state.sim_config_planta["acepta_peligrosos"] = st.checkbox("¿Acepta residuos peligrosos?", value=st.session_state.sim_config_planta["acepta_peligrosos"])
            st.info("💡 Los residuos habilitados se configuran en la tabla de Recepción.")
    
    st.markdown("---")
    
    # 1. RECEPCIÓN
    st.markdown("##### 🚛 1. Recepción y Tipos de Residuo")
    modo_rec = st.radio("Modo de entrada:", ["Agregado (toneladas por tipo)", "Detallado (lista de camiones)"], horizontal=True)
    
    if "Agregado" in modo_rec:
        df_res = pd.DataFrame(st.session_state.sim_tipos_residuo) if st.session_state.sim_tipos_residuo else pd.DataFrame(columns=["Residuo", "Toneladas", "Recuperación (%)", "Rechazo (%)", "Pérdidas (%)"])
        res_ok = st.session_state.sim_config_planta["residuos_habilitados"]
        df_res = df_res[df_res["Residuo"].isin(res_ok)] if not df_res.empty else df_res
        if df_res.empty:
            df_res = pd.DataFrame([{"Residuo": r, "Toneladas": 0.0, "Recuperación (%)": 65.0, "Rechazo (%)": 30.0, "Pérdidas (%)": 5.0} for r in res_ok])
        
        edit_res = st.data_editor(df_res, num_rows="dynamic", use_container_width=True, key="sim_editor_res")
        st.session_state.sim_tipos_residuo = edit_res.to_dict("records")
        
    else:
        st.info("Introduce los camiones manualmente. El sistema sumará el peso neto por tipo de residuo.")
        df_cam = pd.DataFrame(st.session_state.sim_camiones) if st.session_state.sim_camiones else pd.DataFrame(columns=["Matrícula", "Residuo", "Peso Bruto (kg)", "Tara (kg)"])
        edit_cam = st.data_editor(df_cam, num_rows="dynamic", use_container_width=True, key="sim_editor_cam")
        st.session_state.sim_camiones = edit_cam.to_dict("records")
        
        if st.button("🔄 Calcular totales desde camiones"):
            totales = {}
            for c in st.session_state.sim_camiones:
                neto = max(0, c.get("Peso Bruto (kg)", 0) - c.get("Tara (kg)", 0)) / 1000.0
                r = c.get("Residuo", "")
                if r: totales[r] = totales.get(r, 0) + neto
            
            st.session_state.sim_tipos_residuo = []
            for r, t in totales.items():
                st.session_state.sim_tipos_residuo.append({"Residuo": r, "Toneladas": t, "Recuperación (%)": 65.0, "Rechazo (%)": 30.0, "Pérdidas (%)": 5.0})
            st.success("Totales calculados. Revisa las toneladas en la tabla de clasificación.")
            
    residuos_actuales = [t["Residuo"] for t in st.session_state.sim_tipos_residuo if t.get("Toneladas", 0) > 0]
    for r in residuos_actuales:
        if r not in st.session_state.sim_config_planta["residuos_habilitados"]:
            st.warning(f"⚠️ El residuo '{r}' no está habilitado en esta planta.")
        if CATALOGO_RESIDUOS.get(r, {}).get("peligroso") and not st.session_state.sim_config_planta["acepta_peligrosos"]:
            st.error(f"🚨 El residuo '{r}' es PELIGROSO y la planta no está autorizada.")

    st.markdown("---")

    # 2. CLASIFICACIÓN
    st.markdown("##### ⚗️ 2. Clasificación y Matriz de Composición")
    fracciones_necesarias = set()
    for r in st.session_state.sim_config_planta["residuos_habilitados"]:
        if r in CATALOGO_RESIDUOS:
            fracciones_necesarias.update(CATALOGO_RESIDUOS[r]["fracciones"])
            
    df_frac = pd.DataFrame(st.session_state.sim_fracciones) if st.session_state.sim_fracciones else pd.DataFrame(columns=["Fracción", "Precio Venta (€/t)", "% Expedido"])
    if df_frac.empty:
        df_frac = pd.DataFrame([{"Fracción": f, "Precio Venta (€/t)": 100.0, "% Expedido": 80.0} for f in fracciones_necesarias])
    
    edit_frac = st.data_editor(df_frac, num_rows="dynamic", use_container_width=True, key="sim_editor_frac")
    st.session_state.sim_fracciones = edit_frac.to_dict("records")
    
    st.markdown("**Matriz de composición (porcentaje dentro del material recuperado de cada residuo, debe sumar 100%):**")
    matriz_data = []
    for t in st.session_state.sim_tipos_residuo:
        r_nom = t["Residuo"]
        fila = {"Residuo": r_nom}
        for f in st.session_state.sim_fracciones:
            f_nom = f["Fracción"]
            fila[f_nom] = st.session_state.sim_matriz.get(r_nom, {}).get(f_nom, 0)
        matriz_data.append(fila)
        
    if matriz_data:
        df_matriz = pd.DataFrame(matriz_data)
        edit_matriz = st.data_editor(df_matriz, use_container_width=True, key="sim_editor_matriz", disabled=["Residuo"])
        
        nueva_matriz = {}
        errores_matriz = []
        for idx, row in edit_matriz.iterrows():
            r_nom = row["Residuo"]
            nueva_matriz[r_nom] = {}
            suma = 0
            for f in st.session_state.sim_fracciones:
                f_nom = f["Fracción"]
                val = row.get(f_nom, 0)
                nueva_matriz[r_nom][f_nom] = val
                suma += val
            if abs(suma - 100) > 0.1:
                errores_matriz.append(f"'{r_nom}' suma {suma}% (debe ser 100%)")
                
        st.session_state.sim_matriz = nueva_matriz
        
        if errores_matriz:
            st.error("Errores en la matriz: " + ", ".join(errores_matriz))

    st.markdown("---")
    
    # 3. COSTES Y STOCK
    st.markdown("##### 💰 3. Costes y Stock Inicial")
    c1, c2, c3, c4 = st.columns(4)
    with c1: st.session_state.sim_costes["tarifa"] = st.number_input("Tarifa tratamiento (€/t):", value=st.session_state.sim_costes["tarifa"])
    with c2: st.session_state.sim_costes["variable"] = st.number_input("Coste variable (€/t):", value=st.session_state.sim_costes["variable"])
    with c3: st.session_state.sim_costes["fijos"] = st.number_input("Costes fijos (€):", value=st.session_state.sim_costes["fijos"])
    with c4: st.session_state.sim_costes["rechazo"] = st.number_input("Coste rechazo (€/t):", value=st.session_state.sim_costes["rechazo"])

    st.markdown("**Stock Inicial por Fracción (t):**")
    stock_dict = {}
    for f in st.session_state.sim_fracciones:
        f_nom = f["Fracción"]
        val = st.number_input(f"Stock {f_nom}", min_value=0.0, value=float(st.session_state.sim_stock_inicial.get(f_nom, 0)), step=10.0, key=f"stock_{f_nom}")
        stock_dict[f_nom] = val
    st.session_state.sim_stock_inicial = stock_dict

    st.markdown("---")
    
    # 4. RESULTADOS
    if st.button("🚀 EJECUTAR SIMULACIÓN", type="primary", use_container_width=True):
        if not matriz_data: st.error("Añade residuos para simular."); st.stop()
        if errores_matriz: st.error("Corrige la matriz antes de simular."); st.stop()
        
        params = {
            "tipos_residuo": st.session_state.sim_tipos_residuo,
            "fracciones": st.session_state.sim_fracciones,
            "matriz": st.session_state.sim_matriz,
            "costes": st.session_state.sim_costes,
            "stock_inicial": st.session_state.sim_stock_inicial,
            "durabilidades": st.session_state.sim_durabilidades,
            "config_planta": st.session_state.sim_config_planta
        }
        
        res = calcular_simulacion(params)
        if res:
            st.session_state.sim_resultado_actual = res
            st.session_state.sim_params_actuales = params
            st.success("Simulación completada.")
        else:
            st.error("Error en el cálculo.")

    if st.session_state.sim_resultado_actual:
        r = st.session_state.sim_resultado_actual
        st.markdown("### 📊 Resultados de la Simulación")
        
        if r["avisos_durabilidad"]:
            for av in r["avisos_durabilidad"]:
                if av["estado"] == "Vencido": st.error(f"🚨 DURABILIDAD: El material {av['material']} lleva ~{av['edad']} días (max {av['durabilidad']}). ¡Riesgo de degradación!")
                else: st.warning(f"⚠️ DURABILIDAD: El material {av['material']} está cerca de su límite ({av['edad']}/{av['durabilidad']} días).")

        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Total Entrada", f"{r['total_ton']:,.1f} t")
        c2.metric("Ingresos", f"{r['ingresos_totales']:,.0f} €")
        c3.metric("Costes", f"{r['costes_totales']:,.0f} €")
        c4.metric("Beneficio", f"{r['beneficio']:,.0f} €", delta=f"{r['beneficio']/r['total_ton']:.1f} €/t" if r['total_ton'] else "0")
        
        st.markdown("---")
        st.markdown("#### 💾 Guardar este escenario")
        c_g1, c_g2 = st.columns(2)
        with c_g1: nom_esc = st.text_input("Nombre del escenario:")
        with c_g2: cli_esc = st.text_input("Cliente / Empresa (opcional):")
        
        if st.button("💾 Guardar Escenario en Base de Datos", type="primary"):
            if nom_esc:
                ref = f"SIM-{datetime.datetime.now().strftime('%Y%m%d-%H%M%S')}"
                insert_db(
                    "INSERT INTO escenarios_simulacion (fecha, referencia, nombre, cliente, notas, parametros_json, ingresos, costes, beneficio) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)",
                    (datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), ref, nom_esc, cli_esc, "", json.dumps(st.session_state.sim_params_actuales), r['ingresos_totales'], r['costes_totales'], r['beneficio'])
                )
                st.success(f"✅ Escenario '{nom_esc}' guardado con referencia {ref}.")
            else:
                st.error("Pon un nombre al escenario.")

def render_inventario():
    st.markdown("#### 📦 Inventario Vivo y Expediciones")
    if not st.session_state.sim_resultado_actual:
        st.info("Ejecuta primero una simulación en la pestaña 🧮 SIMULAR.")
        return
        
    r = st.session_state.sim_resultado_actual
    st.markdown("##### Estado Final del Almacén")
    
    rows = []
    for f in st.session_state.sim_fracciones:
        f_nom = f["Fracción"]
        s_ini = st.session_state.sim_stock_inicial.get(f_nom, 0)
        prod = r["produccion"].get(f_nom, 0)
        exp = r["expedido"].get(f_nom, 0)
        fin = r["stock_final"].get(f_nom, 0)
        rows.append({"Fracción": f_nom, "Stock Inicial": s_ini, "Producido": prod, "Expedido": exp, "Stock Final": fin})
        
    st.dataframe(pd.DataFrame(rows).round(1), use_container_width=True, hide_index=True)
    
    if r["avisos_durabilidad"]:
        st.markdown("##### 🚨 Avisos de Durabilidad")
        for av in r["avisos_durabilidad"]:
            st.warning(f"Material **{av['material']}**: {av['estado']} (Edad est. {av['edad']} d / Max {av['durabilidad']} d)")

def render_mantenimiento_legal():
    st.markdown("#### 🔧 Mantenimiento Legal")
    st.caption("Frecuencias orientativas. Validar con la OCA y la AAI específica de la planta.")
    
    df = get_db_data("mantenimiento_legal")
    if df.empty:
        if st.button("🎲 Cargar biblioteca de tareas legales"):
            for t in BIBLIOTECA_MANTENIMIENTO:
                prox = (datetime.datetime.now() + datetime.timedelta(days=t["frecuencia_dias"])).strftime("%Y-%m-%d")
                insert_db("INSERT INTO mantenimiento_legal (planta, equipo, norma, frecuencia_dias, ultima_ejecucion, proxima_ejecucion, estado) VALUES (?, ?, ?, ?, ?, ?, ?)",
                          ("Planta Demo", t["equipo"], t["norma"], t["frecuencia_dias"], datetime.datetime.now().strftime("%Y-%m-%d"), prox, "OK"))
            st.rerun()
    else:
        st.dataframe(df[["equipo", "norma", "frecuencia_dias", "proxima_ejecucion", "estado"]], use_container_width=True)

def render_comparar():
    st.markdown("#### 📊 Comparar Escenarios Guardados")
    df = get_db_data("escenarios_simulacion")
    if df.empty:
        st.info("No hay escenarios guardados.")
        return
        
    st.dataframe(df[["id", "fecha", "nombre", "cliente", "ingresos", "costes", "beneficio"]], use_container_width=True, hide_index=True)
    
    fig = px.bar(df, x="nombre", y=["ingresos", "costes", "beneficio"], barmode="group", title="Comparativa Económica")
    st.plotly_chart(fig, use_container_width=True)

def render_informe():
    st.markdown("#### 📄 Generar Informe Profesional")
    df = get_db_data("escenarios_simulacion")
    if df.empty:
        st.info("No hay escenarios guardados para generar un informe.")
        return
        
    opciones = {f"#{row['id']} - {row['nombre']} ({row['fecha']})": row for _, row in df.iterrows()}
    sel = st.selectbox("Selecciona un escenario:", list(opciones.keys()))
    
    if st.button("📄 GENERAR INFORME (PDF/Word)", type="primary"):
        row = opciones[sel]
        params = json.loads(row["parametros_json"])
        
        with st.spinner("Redactando informe con IA..."):
            prompt_ia = f"Redacta un resumen ejecutivo para un informe de planta de residuos. Datos: {row['nombre']}. Ingresos: {row['ingresos']}€, Costes: {row['costes']}€, Beneficio: {row['beneficio']}€."
            resumen_ia = llamar_ia("llama3", prompt_ia)
            
        texto_informe = f"""
INFORME DE SIMULACIÓN DE PLANTA DE RESIDUOS
===========================================
Referencia: {row['referencia']}
Cliente: {row['cliente']}
Fecha: {row['fecha']}

1. RESUMEN EJECUTIVO
--------------------
{resumen_ia}

2. DATOS ECONÓMICOS
-------------------
Ingresos Totales: {row['ingresos']:,.2f} €
Costes Totales:   {row['costes']:,.2f} €
Beneficio:        {row['beneficio']:,.2f} €

3. PARÁMETROS DE SIMULACIÓN
---------------------------
(Ver anexo JSON)

Aviso legal: Documento generado a partir de simulación con datos hipotéticos.
"""
        botones_descarga(texto_informe, f"Informe_{row['nombre']}")
# ==========================================
# INTERFAZ PRINCIPAL Y CONTROL DE MENÚ
# ==========================================
if not st.session_state.sidebar_visible:
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        if st.button("⛶ MOSTRAR MENÚ Y KEYS", type="primary", use_container_width=True):
            st.session_state.sidebar_visible = True
            st.rerun()

st.markdown("<h1>XAVILU NETWORK</h1><p style='text-align:center; color:#FFD700; font-weight:bold; letter-spacing:2px; text-shadow: 0 0 10px #FFD700;'>● SISTEMA CONECTADO</p>", unsafe_allow_html=True)

if st.session_state.sidebar_visible:
    with st.sidebar:
        try: st.image("logo.png", use_container_width=True)
        except: st.markdown("### 🪐 PANEL DE CONTROL")
        
        if st.button("⛶ Ocultar Menú (Pantalla Completa)", use_container_width=True): 
            st.session_state.sidebar_visible = False
            st.rerun()
            
        if st.button("🚪 Cerrar Sesión", use_container_width=True):
            st.session_state["password_correct"] = False
            st.rerun()
            
        st.markdown("---")
        modelo_seleccionado = st.selectbox("Ollama Local", ["llama3.2:latest", "qwen2.5:7b", "deepseek-r1:1.5b", "llama3.2-vision:latest"])
        
        st.markdown("#### 🔑 Estado de Conexiones")
        try:
            if st.secrets.get("GEMINI_API_KEY"): st.success("✅ Gemini Conectado (vía TOML)")
        except: pass
        
        if st.session_state.get("rol") == "admin":
            with st.expander("⚙️ Configurar Keys Manualmente"):
                st.caption("Si no usas secrets.toml, pega aquí tus claves para esta sesión:")
                st.session_state.api_key_gemini_manual = st.text_input("Gemini API Key", type="password", value=st.session_state.api_key_gemini_manual)
                st.session_state.api_key_groq_manual = st.text_input("Groq API Key", type="password", value=st.session_state.api_key_groq_manual)
                st.session_state.api_key_silicon_manual = st.text_input("SiliconFlow API Key", type="password", value=st.session_state.api_key_silicon_manual)
                st.session_state.api_key_mistral_manual = st.text_input("Mistral API Key", type="password", value=st.session_state.api_key_mistral_manual)
                st.session_state.api_key_openrouter_manual = st.text_input("OpenRouter API Key", type="password", value=st.session_state.api_key_openrouter_manual)

        st.markdown("---")
        st.session_state.whatsapp_number = st.text_input("📲 WhatsApp (ej. 34600000000):", value=st.session_state.whatsapp_number)
        st.markdown("---")
        archivo_global = st.file_uploader("📁 Cerebro Global", type=["pdf", "txt", "docx", "png", "jpg"], key="up_global")
        if archivo_global and st.session_state.get("global_file_name") != archivo_global.name:
            with st.spinner("Procesando..."):
                st.session_state.global_context = leer_archivo_universal(archivo_global, modelo_seleccionado)[:5000]
                st.session_state.global_file_name = archivo_global.name; st.success("✅ Memorizado.")
        if st.session_state.get("global_file_name"):
            st.info(f"Activo: **{st.session_state.global_file_name}**")
            if st.button("🗑️ Olvidar"): st.session_state.global_context = ""; st.session_state.global_file_name = ""; st.rerun()
        st.markdown("---")
        
        if st.session_state.get("rol") == "admin":
            st.markdown("#### 📂 Historial Reciente")
            for h in obtener_de_db():
                with st.expander(f"{h[1]} - {h[0][:10]}"): 
                    st.caption(h[2][:200] + "...")
                    if st.button("Recuperar", key=f"rec_{h[0]}_{h[1]}"):
                        st.session_state.office_borrador = h[2]
                        st.session_state.podcast_script = h[2]
                        st.success("¡Texto cargado en el borrador del Podcast/Office!")
else:
    modelo_seleccionado = "llama3.2:latest"
    api_key_gemini = st.session_state.api_key_gemini_manual

# MENÚ PRINCIPAL
macro_names = ["🛠️ XAVIWORK", "💬 CHAT Y ASISTENTES", "📁 DOCUMENTOS Y TEXTO", "🌐 WEB Y ACTUALIDAD", "🎬 MULTIMEDIA", "🎓 ESTUDIO Y OFICIOS", "🌍 VIAJES Y HOGAR", "♻️ GESTIÓN PLANTAS", "🎵 MÚSICA Y ARTE"]
macro_idx = custom_tabs(macro_names, "macro", 0)
st.markdown("<div style='margin-bottom: 15px;'></div>", unsafe_allow_html=True)

# ==========================================
# CATEGORÍA 0: XAVIWORK
# ==========================================
if macro_idx == 0:
    st.markdown("#### 🛠️ XaviWork - El Estudio Definitivo")
    contexto = cajita_adjuntar_archivo("xaviwork_adj", modelo_seleccionado, st.session_state.api_key_gemini_manual)
    
    tema_work = campo_con_ideas("¿Sobre qué quieres trabajar?", "xaviwork", "estrategias y proyectos", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Estrategia de marketing...", "", "")
    
    if st.button("✨ Generar Propuesta Final de XaviWork", type="primary"):
        if tema_work:
            with st.spinner("Trabajando en tu proyecto..."):
                prompt = f"{contexto}\n\nActúa como un experto consultor. Desarrolla la siguiente idea de forma detallada y estructurada:\n{tema_work}"
                res = llamar_ia(modelo_seleccionado, prompt, st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else:
                    st.session_state.work_result = res
                    guardar_en_db("XaviWork", res)
        else:
            st.warning("Escribe sobre qué quieres trabajar primero.")

    if st.session_state.get("work_result"):
        st.markdown("### 📝 Resultado de XaviWork")
        texto_editado = st.text_area("Modifica el resultado:", value=st.session_state.work_result, height=400, key="xaviwork_txt")
        st.session_state.work_result = texto_editado
        botones_descarga(texto_editado, "XaviWork_Proyecto")

# ==========================================
# CATEGORÍA 1: CHAT Y ASISTENTES IA
# ==========================================
elif macro_idx == 1:
    sub_names = ["🎙️ PODCAST", "💬 CHAT", "🎙️ VOZ", "💬 CHAT CON PDF", "🕵️‍♂️ AGENTES", "⚖️ DEBATE", "🍳 COCINA", "📖 CUENTOS", "💻 CÓDIGO", "🌍 TRADUCTOR"]
    sub_idx = custom_tabs(sub_names, "chat_sub", 0)
    
    if sub_idx == 0:
        st.markdown("#### 🎙️ Generador de Podcasts")
        contexto = cajita_adjuntar_archivo("podcast_adj", modelo_seleccionado, st.session_state.api_key_gemini_manual)
        tema = st.text_area("Temática principal:", placeholder="Ej: El futuro de la IA, La pizza con piña...", key="pod_tema_input", height=100)
        
        c1, c2 = st.columns(2)
        with c1: estilo = st.selectbox("Estilo:", ["Entrevista", "Debate", "Educativo", "Narrativo"], key="p_est")
        with c2: dur = st.slider("Duración:", 1, 10, 3, key="p_dur")
        
        st.markdown("---")
        if st.button("💡 Generar 3 Ideas con IA (basadas en tu tema)", use_container_width=True, key="pod_btn_ideas"):
            with st.spinner("Pensando ideas..."):
                prompt_ideas = f"Dame 3 ideas creativas y concisas para un podcast de {dur} minutos, estilo {estilo}, sobre el tema: '{tema}'. Devuelve SOLO las 3 ideas separadas por saltos de línea, sin numeración."
                res = llamar_ia(modelo_seleccionado, prompt_ideas, st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                st.session_state["pod_sugs"] = [l.strip("-*1234567890. ") for l in res.split('\n') if len(l.strip()) > 5][:3]

        if st.session_state.get("pod_sugs"):
            st.markdown("✨ **Selecciona las ideas que quieras añadir al tema (es opcional):**")
            seleccionadas_pod = []
            for i, idea in enumerate(st.session_state["pod_sugs"]):
                if st.checkbox(idea, key=f"pod_chk_{i}"):
                    seleccionadas_pod.append(idea)

            c_pod1, c_pod2 = st.columns(2)
            with c_pod1:
                if st.button("➕ Añadir ideas seleccionadas al tema", key="pod_add_selected", use_container_width=True, type="primary"):
                    if seleccionadas_pod:
                        texto_actual = st.session_state["pod_tema_input"].strip()
                        nuevo_texto = "\n\n".join(seleccionadas_pod)
                        if texto_actual:
                            st.session_state["pod_tema_input"] = f"{texto_actual}\n\n{nuevo_texto}"
                        else:
                            st.session_state["pod_tema_input"] = nuevo_texto
                    for j in range(len(st.session_state["pod_sugs"])):
                        k = f"pod_chk_{j}"
                        if k in st.session_state: del st.session_state[k]
                    del st.session_state["pod_sugs"]
                    st.rerun()
            with c_pod2:
                if st.button("🗑️ Descartar ideas", key="pod_discard", use_container_width=True):
                    for j in range(len(st.session_state["pod_sugs"])):
                        k = f"pod_chk_{j}"
                        if k in st.session_state: del st.session_state[k]
                    del st.session_state["pod_sugs"]
                    st.rerun()

        st.markdown("---")
        if st.button("✨ Crear Guion", type="primary", key="btn_pod") and st.session_state["pod_tema_input"]:
            with st.spinner("Escribiendo..."):
                prompt = f"{contexto}\n\nGuion podcast {dur} min. Tema: '{st.session_state['pod_tema_input']}'. Estilo: {estilo}. 3 locutores. REGLA ESTRICTA: Escribe de forma muy natural y coloquial. NO uses acotaciones como (risas) o (pausa). Si quieres que se rían, escribe 'jajaja'. Si quieres que piensen, escribe 'mmm'. Usa SOLO el formato exacto: 'Locutor 1: [texto]', 'Locutor 2: [texto]', 'Locutor 3: [texto]'."
                res = llamar_ia(modelo_seleccionado, prompt, st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.session_state.podcast_script = res; guardar_en_db("Podcast", res)

        if st.session_state.get("podcast_script"):
            st.markdown("### 📜 Guion Generado")
            st.text_area("Guion (Editable):", value=st.session_state.podcast_script, height=300, key="pod_edit")
            st.session_state.podcast_script = st.session_state.pod_edit
            c1, c2, c3 = st.columns(3)
            with c1: v1 = st.selectbox("Voz L1:", ["es-ES-ElviraNeural", "es-MX-JorgeNeural"], key="v1")
            with c2: v2 = st.selectbox("Voz L2:", ["es-ES-ElviraNeural", "es-MX-JorgeNeural"], index=1, key="v2")
            with c3: v3 = st.selectbox("Voz L3:", ["es-AR-TomasNeural", "es-ES-ElviraNeural"], key="v3")
            
            c_v1, c_v2 = st.columns(2)
            with c_v1:
                if st.button("🔊 Audio 3 Voces", key="btn_pod_aud"):
                    with st.spinner("Generando audio..."):
                        aud = generar_podcast_3_voces(st.session_state.podcast_script, v1, v2, v3)
                        if aud: st.audio(aud, format="audio/mp3"); st.download_button("💾 MP3", aud, "podcast.mp3")
                        else: st.error("❌ Error al generar el audio.")
            with c_v2:
                if st.button("🎨 Generar Portada", key="btn_pod_cover", type="secondary"):
                    with st.spinner("Diseñando portada..."):
                        tema_portada = st.session_state.get('pod_tema_input', 'Podcast')[:50]
                        url_img = f"https://image.pollinations.ai/prompt/{urllib.parse.quote('Podcast cover about ' + tema_portada + ', high quality')}?width=1024&height=1024&nologo=true"
                        st.image(url_img, caption="Portada generada por IA")
                        st.session_state.galeria_imgs.append(url_img)
            
            botones_descarga(st.session_state.podcast_script, "Podcast")

    elif sub_idx == 1:
        st.markdown("#### 💬 Chat")
        for m in st.session_state.chat_history: st.chat_message(m["role"]).write(m["content"])
        if st.session_state.chat_history:
            c1, c2, c3 = st.columns(3)
            with c1: st.download_button("📥 TXT", "\n".join([f"{m['role']}: {m['content']}" for m in st.session_state.chat_history]), "chat.txt")
            with c2: 
                if st.button("🗑️ Limpiar"): st.session_state.chat_history = []; limpiar_chat_db(); st.rerun()
            with c3:
                if st.button("🔊 Leer"):
                    aud = generar_audio_sincrono(st.session_state.chat_history[-1]["content"], "es-ES-ElviraNeural")
                    if aud: st.audio(aud, format="audio/mp3")
        
        ctx = cajita_adjuntar_archivo("chat_adj", modelo_seleccionado, st.session_state.api_key_gemini_manual)
        
        c_s1, c_s2, c_s3, c_s4 = st.columns(4)
        accion_rapida = None
        if c_s1.button("Resumir", key="q_res"): accion_rapida = "Resume esto de forma concisa:"
        if c_s2.button("Traducir", key="q_trad"): accion_rapida = "Traduce al inglés:"
        if c_s3.button("Formal", key="q_form"): accion_rapida = "Hazlo más formal y profesional:"
        if c_s4.button("Buscar Web", key="q_web"): accion_rapida = "Busca en internet sobre esto y resume:"

        st.markdown("##### 🎤 Dictado por voz (Opcional)")
        audio_value = st.audio_input("Graba tu mensaje aquí", key="voice_input")
        
        p = chat_input_pro("main")
        
        if audio_value:
            with st.spinner("Escuchando y transcribiendo..."):
                transcripcion = analizar_archivo_con_gemini(audio_value.getvalue(), "audio/wav", "Transcribe exactamente lo que dice este audio. Devuelve solo el texto.", st.session_state.api_key_gemini_manual)
                if transcripcion and "[ERROR" not in transcripcion:
                    p = transcripcion.strip()

        if p:
            prompt_chat = f"{accion_rapida}\n\n{p}" if accion_rapida else p
            st.session_state.chat_history.append({"role": "user", "content": p}); st.chat_message("user").write(p)
            guardar_chat_db("user", p)
            with st.spinner("Pensando..."):
                res = llamar_ia(modelo_seleccionado, f"{ctx}\n\n{prompt_chat}", st.session_state.api_key_gemini_manual, st.session_state.chat_history, "", "")
                if "[ERROR" in res: st.error(res)
                else:
                    st.session_state.chat_history.append({"role": "assistant", "content": res}); st.chat_message("assistant").write(res); guardar_en_db("Chat", res); guardar_chat_db("assistant", res)
            st.rerun()

    elif sub_idx == 2:
        st.markdown("#### Voz")
        txt = st.text_area("Texto:")
        voz = st.selectbox("Voz:", ["es-ES-ElviraNeural", "es-MX-JorgeNeural"])
        if st.button("🔊 Audio") and txt:
            aud = generar_audio_sincrono(txt, voz)
            if aud: st.audio(aud, format="audio/mp3")

    elif sub_idx == 3:
        st.markdown("#### Chat con PDFs")
        archs = st.file_uploader("Sube docs:", type=["pdf", "txt", "docx"], accept_multiple_files=True)
        if archs:
            texto = "".join([leer_archivo_universal(a, modelo_seleccionado) for a in archs])[:MAX_TEXTO_IA]
            st.success("Docs cargados.")
            for m in st.session_state.pdf_chat_history: st.chat_message(m["role"]).write(m["content"])
            p = chat_input_pro("pdf")
            if p:
                st.session_state.pdf_chat_history.append({"role": "user", "content": p})
                with st.spinner("Buscando..."):
                    res = llamar_ia(modelo_seleccionado, f"Contexto:\n{texto}\n\nPregunta: {p}", st.session_state.api_key_gemini_manual, st.session_state.pdf_chat_history, "", "")
                    if "[ERROR" in res: st.error(res)
                    else: st.session_state.pdf_chat_history.append({"role": "assistant", "content": res}); st.rerun()

    elif sub_idx == 4:
        st.markdown("#### Agentes")
        nom = st.text_input("Nombre:"); rol = st.text_area("Personalidad:")
        if st.button("Crear"): st.session_state.agente_activo = {"nombre": nom, "rol": rol}; st.session_state.agente_chat_history = []
        if st.session_state.get("agente_activo"):
            ag = st.session_state.agente_activo
            for m in st.session_state.agente_chat_history: st.chat_message(m["role"]).write(m["content"])
            p = chat_input_pro("agent")
            if p:
                st.session_state.agente_chat_history.append({"role": "user", "content": p})
                with st.spinner("Pensando..."):
                    res = llamar_ia(modelo_seleccionado, f"ERES: {ag['rol']}\nUsuario: {p}", st.session_state.api_key_gemini_manual, st.session_state.agente_chat_history, "", "")
                    if "[ERROR" in res: st.error(res)
                    else: st.session_state.agente_chat_history.append({"role": "assistant", "content": res}); st.rerun()

    elif sub_idx == 5:
        st.markdown("#### Debate")
        t = campo_con_ideas("Tema:", "deb", "temas de debate", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Coche eléctrico", "", "")
        if st.button("Iniciar") and t:
            res = llamar_ia(modelo_seleccionado, f"Debate {t}: 3 a favor, 3 en contra, conclusión.", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.markdown(res); botones_descarga(res, "Debate")

    elif sub_idx == 6:
        st.markdown("#### Cocina")
        ing = campo_con_ideas("Ingredientes:", "chef1", "recetas", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Pollo", "", "")
        pref = campo_con_ideas("Menú:", "chef2", "menús", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Vegetariano", "", "")
        c1, c2 = st.columns(2)
        with c1:
            if st.button("Receta") and ing: 
                res = llamar_ia(modelo_seleccionado, f"Receta con {ing}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.info(res); botones_descarga(res, "Receta")
        with c2:
            if st.button("Menú") and pref: 
                res = llamar_ia(modelo_seleccionado, f"Menú semanal {pref}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.info(res); botones_descarga(res, "Menu")

    elif sub_idx == 7:
        st.markdown("#### Cuentos")
        p = campo_con_ideas("Protagonista:", "c1", "protagonistas", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Dragón", "", "")
        l = campo_con_ideas("Lugar:", "c2", "lugares mágicos", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Bosque", "", "")
        if st.button("Crear") and p and l:
            res = llamar_ia(modelo_seleccionado, f"Cuento sobre {p} en {l}.", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else:
                st.text_area("Cuento:", res, height=300); botones_descarga(res, "Cuento")
                if st.button("Leer"): 
                    aud = generar_audio_sincrono(res, "es-ES-ElviraNeural")
                    if aud: st.audio(aud, format="audio/mp3")

    elif sub_idx == 8:
        st.markdown("#### Código")
        req = campo_con_ideas("Código:", "code", "scripts", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Renombrar archivos", "", "")
        if st.button("Procesar") and req:
            res = llamar_ia(modelo_seleccionado, f"Programador senior. Explica o mejora:\n{req}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.code(res, language="python"); botones_descarga(res, "Codigo")

    elif sub_idx == 9:
        st.markdown("#### Traductor")
        txt = campo_con_ideas("Texto:", "trad", "frases", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Texto...", "", "")
        c1, c2 = st.columns(2)
        with c1: o = st.selectbox("Origen:", ["Auto", "Español", "Inglés"], key="o")
        with c2: d = st.selectbox("Destino:", ["Español", "Inglés", "Francés"], key="d")
        if st.button("Traducir") and txt:
            res = llamar_ia(modelo_seleccionado, f"Traduce del {o} al {d}:\n{txt}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.info(res); botones_descarga(res, "Traduccion")

# ==========================================
# CATEGORÍA 2: DOCUMENTOS Y TEXTO
# ==========================================
elif macro_idx == 2:
    sub_names = ["📁 ANÁLISIS", "⚖️ COMPARAR", "📊 OFFICE", "📄 PDF", "🕸️ MAPAS", "📇 QR & CV", "📋 PORTAPAPELES"]
    sub_idx = custom_tabs(sub_names, "doc_sub", 0)
    
    if sub_idx == 0:
        st.markdown("#### Análisis de Documentos")
        arch = st.file_uploader("Sube PDF, Word, Excel o TXT", type=["pdf", "txt", "docx", "xlsx"])
        opc = st.selectbox("Análisis:", ["Resumen Ejecutivo", "Extraer Fechas y Cifras", "Buscar Errores", "Crear Test", "Traducir al Inglés"])
        if arch and st.button("🔎 Ejecutar", type="primary"):
            with st.spinner("Analizando..."):
                texto = leer_archivo_universal(arch, modelo_seleccionado)
                prompt = f"Resumen:\n{texto}" if opc == "Resumen Ejecutivo" else f"Extrae fechas:\n{texto}" if opc == "Extraer Fechas y Cifras" else f"Errores:\n{texto}" if opc == "Buscar Errores" else f"Test 5 preguntas:\n{texto}" if opc == "Crear Test" else f"Traduce al inglés:\n{texto}"
                res = llamar_ia(modelo_seleccionado, prompt, st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.text_area("Resultado:", res, height=400); botones_descarga(res, "Analisis")

    elif sub_idx == 1:
        st.markdown("#### Comparador")
        a = st.file_uploader("Original:", type=["pdf", "txt"]); b = st.file_uploader("Modificado:", type=["pdf", "txt"])
        if a and b and st.button("Comparar"):
            with st.spinner("Buscando..."):
                ta = leer_archivo_universal(a, modelo_seleccionado); tb = leer_archivo_universal(b, modelo_seleccionado)
                res = llamar_ia(modelo_seleccionado, f"Diferencias:\nA:\n{ta[:10000]}\nB:\n{tb[:10000]}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.text_area("Diferencias:", res, height=400); botones_descarga(res, "Comparativa")

    elif sub_idx == 2:
        st.markdown("#### Creación de Documentos (Flujo Borrador)")
        contexto = cajita_adjuntar_archivo("office_adj", modelo_seleccionado, st.session_state.api_key_gemini_manual)
        
        st.markdown("**Elige el estilo de la IA:**")
        c_est1, c_est2, c_est3 = st.columns(3)
        if c_est1.button("👔 Profesional", use_container_width=True): st.session_state.office_estilo = "Profesional y formal"
        if c_est2.button("🎨 Creativo", use_container_width=True): st.session_state.office_estilo = "Creativo y original"
        if c_est3.button("👨‍💻 Técnico", use_container_width=True): st.session_state.office_estilo = "Técnico y preciso"

        prompt_office = campo_con_ideas("¿Qué documento quieres crear?", "office", "documentos", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Presupuesto para cliente", "", "")
        formato = st.radio("Formato final:", ["Word", "Excel", "PowerPoint"], key="office_fmt", horizontal=True)
        
        if st.button("✨ Generar Borrador de Texto", type="primary", key="btn_office_gen"):
            if prompt_office:
                with st.spinner("Generando borrador..."):
                    prompt_ia = f"Actúa con un estilo {st.session_state.office_estilo}.\n\n{contexto}\n\n{prompt_office}"
                    if formato == "PowerPoint": prompt_ia += "\n\nIMPORTANTE: Formatea para diapositivas. Título, salto de línea, guiones para los puntos. Doble salto de línea entre diapositivas."
                    elif formato == "Excel": prompt_ia += "\n\nIMPORTANTE: Devuelve los datos en formato CSV separados por comas."
                    res = llamar_ia(modelo_seleccionado, prompt_ia, st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                    if "[ERROR" in res: st.error(res)
                    else: st.session_state.office_borrador = res
            else: st.warning("Escribe un prompt primero.")

        if st.session_state.get("office_borrador"):
            st.markdown("### 📝 Borrador Editable")
            texto_editado = st.text_area("Modifica tu documento antes de descargar:", value=st.session_state.office_borrador, height=350, key="office_txt")
            st.session_state.office_borrador = texto_editado
            
            c1, c2, c3 = st.columns(3)
            with c1: boton_superpoder_borrador(texto_editado, "Hazlo más profesional", "of_sup1", modelo_seleccionado, st.session_state.api_key_gemini_manual, "office_borrador", "", "")
            with c2: boton_superpoder_borrador(texto_editado, "Resúmelo", "of_sup2", modelo_seleccionado, st.session_state.api_key_gemini_manual, "office_borrador", "", "")
            with c3: boton_superpoder_borrador(texto_editado, "Añade más detalles", "of_sup3", modelo_seleccionado, st.session_state.api_key_gemini_manual, "office_borrador", "", "")

            st.markdown("---")
            botones_descarga(texto_editado, "Documento")

    elif sub_idx == 3:
        st.markdown("#### PDF Pro (Flujo Borrador)")
        contexto = cajita_adjuntar_archivo("pdf_adj", modelo_seleccionado, st.session_state.api_key_gemini_manual)
        titulo = st.text_input("Título del PDF:")
        contenido = campo_con_ideas("Contenido:", "pdf_c", "contenido PDF", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Escribe el contenido...", "", "")
        
        if st.button("✨ Generar Borrador PDF", type="primary", key="btn_pdf_gen") and titulo and contenido:
            with st.spinner("Generando..."):
                res = llamar_ia(modelo_seleccionado, f"{contexto}\n\n{contenido}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.session_state.pdf_borrador = res

        if st.session_state.get("pdf_borrador"):
            st.markdown("### 📝 Borrador Editable")
            texto_editado = st.text_area("Modifica antes de descargar:", value=st.session_state.pdf_borrador, height=350, key="pdf_txt")
            st.session_state.pdf_borrador = texto_editado
            
            c1, c2 = st.columns(2)
            with c1: boton_superpoder_borrador(texto_editado, "Hazlo más técnico", "pdf_sup1", modelo_seleccionado, st.session_state.api_key_gemini_manual, "pdf_borrador", "", "")
            with c2: boton_superpoder_borrador(texto_editado, "Añade conclusiones", "pdf_sup2", modelo_seleccionado, st.session_state.api_key_gemini_manual, "pdf_borrador", "", "")

            st.markdown("---")
            botones_descarga(texto_editado, titulo)

    elif sub_idx == 4:
        st.markdown("#### Mapas Mentales")
        t = campo_con_ideas("Tema:", "map", "mapas", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Redes", "", "")
        if st.button("Trazar") and t:
            res = llamar_ia(modelo_seleccionado, f"Genera código Mermaid para mapa mental de '{t}'.", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else:
                st.session_state.mermaid_code = html.escape(res.replace("```mermaid", "").replace("```", "").strip())
                if st.session_state.mermaid_code: st.html(f"<div class='mermaid'>{st.session_state.mermaid_code}</div><script type='module'>import mermaid from 'https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.esm.min.mjs'; mermaid.initialize({{ startOnLoad: true }});</script>")
                botones_descarga(res, "Mapa")

    elif sub_idx == 5:
        c1, c2 = st.columns(2)
        with c1:
            txt = st.text_input("Texto/URL QR:")
            if st.button("Generar QR") and txt:
                qr = qrcode.QRCode(); qr.add_data(txt); qr.make(fit=True)
                buf = BytesIO(); qr.make_image().save(buf, format="PNG"); st.image(buf.getvalue())
        with c2:
            img = st.file_uploader("Leer QR:", type=["png", "jpg"])
            if img and HAS_CV2:
                d, _, _ = cv2.QRCodeDetector().detectAndDecode(np.array(Image.open(img).convert('RGB')))
                if d: st.success(d)

    elif sub_idx == 6:
        st.markdown("#### Portapapeles")
        txt = st.text_area("Pega texto:")
        inst = campo_con_ideas("¿Qué hago?", "clip", "instrucciones", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Resume", "", "")
        if st.button("Procesar") and txt and inst:
            res = llamar_ia(modelo_seleccionado, f"{inst}:\n{txt}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.info(res); botones_descarga(res, "Portapapeles")

# ==========================================
# CATEGORÍA 3: WEB Y ACTUALIDAD
# ==========================================
elif macro_idx == 3:
    sub_names = ["🌐 BUSCADOR", "⚽ DEPORTES", "📺 TV", "📊 WEB A EXCEL", "🎥 YOUTUBE", "📰 LECTOR"]
    sub_idx = custom_tabs(sub_names, "web_sub", 0)
    if sub_idx == 0:
        q = campo_con_ideas("Buscar:", "web", "búsquedas", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Noticias IA", "", "")
        if st.button("🔍 Buscar") and q:
            for r in buscar_en_internet(q): st.markdown(f"**[{r['titulo']}]({r['url']})**"); st.caption(r['resumen'])
    elif sub_idx == 1:
        d = st.selectbox("Deporte:", ["Fútbol", "Baloncesto", "Tenis"])
        if st.button("Noticias"):
            for n in buscar_noticias_deportivas(d, "Todas"): st.markdown(f"**{n['titulo']}** - {n['resumen']}")
    elif sub_idx == 2:
        c = campo_con_ideas("Canal:", "tv", "TV", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Telecinco", "", "")
        if st.button("Programación") and c:
            res = buscar_en_internet(c)
            r_ia = llamar_ia(modelo_seleccionado, f"Programación {c}: {res}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in r_ia: st.error(r_ia)
            else: st.info(r_ia); botones_descarga(r_ia, "TV")
    elif sub_idx == 3:
        url = st.text_input("URL Tabla:")
        if st.button("Extraer") and url:
            try:
                df = pd.read_html(url)[0]; st.dataframe(df)
                out = BytesIO(); df.to_excel(out); st.download_button("📥 Excel", out.getvalue(), "tabla.xlsx")
            except: st.error("Fallo al leer.")
    elif sub_idx == 4:
        url = st.text_input("URL YouTube:")
        if st.button("Resumir") and url:
            try:
                id = url.split("v=")[-1].split("&")[0]; tr = YouTubeTranscriptApi.get_transcript(id, languages=['es', 'en'])
                res = llamar_ia(modelo_seleccionado, f"Resume: {' '.join([t['text'] for t in tr])[:5000]}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.text_area("Resumen:", res, height=200); botones_descarga(res, "Youtube")
            except: st.error("Error.")
    elif sub_idx == 5:
        url = st.text_input("URL Web:")
        if st.button("Leer") and url:
            try:
                txt = " ".join(re.sub('<[^<]+?>', ' ', requests.get(url, headers={'User-Agent': 'Mozilla/5.0'}).text).split())[:5000]
                res = llamar_ia(modelo_seleccionado, f"Resume: {txt}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
                if "[ERROR" in res: st.error(res)
                else: st.text_area("Resumen:", res, height=200); botones_descarga(res, "Web")
            except: st.error("Error.")

# ==========================================
# CATEGORÍA 4: MULTIMEDIA
# ==========================================
elif macro_idx == 4:
    sub_names = ["🖼️ ANÁLISIS", "🎨 VISUAL", "🛡️ LOGOS", "📥 DESCARGADOR", "🖼️ GALERÍA"]
    sub_idx = custom_tabs(sub_names, "multi_sub", 0)
    if sub_idx == 0:
        a = st.file_uploader("Sube archivo:", type=["png", "jpg", "mp3", "wav", "mp4"])
        if a:
            n = a.name.lower()
            if n.endswith((".png", ".jpg")):
                st.image(a); op = st.radio("Opción:", ["Describir", "OCR", "Prompt", "¿Es IA?"])
                if st.button("Ejecutar"):
                    p = "Describe" if "Describir" in op else "OCR" if "OCR" in op else "Prompt Midjourney" if "Prompt" in op else "¿Es IA?"
                    st.info(analizar_archivo_con_gemini(a.getvalue(), "image/jpeg", p, st.session_state.api_key_gemini_manual))
            elif n.endswith((".mp3", ".wav")):
                st.audio(a); op = st.radio("Opción:", ["Transcribir", "Resumen", "Traducir"])
                if st.button("Ejecutar"):
                    p = "Transcribe" if "Transcribir" in op else "Resume" if "Resumen" in op else "Traduce"
                    st.info(analizar_archivo_con_gemini(a.getvalue(), "audio/mp3", p, st.session_state.api_key_gemini_manual))
            elif n.endswith((".mp4")):
                st.video(a); op = st.radio("Opción:", ["Resumir", "Escenas"])
                if st.button("Ejecutar"):
                    p = "Resume" if "Resumir" in op else "Escenas"
                    st.info(analizar_archivo_con_gemini(a.getvalue(), "video/mp4", p, st.session_state.api_key_gemini_manual))
    elif sub_idx == 1:
        p = campo_con_ideas("Describe imagen:", "img", "imágenes", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Gato astronauta", "", "")
        if st.button("Generar") and p:
            with st.spinner("Traduciendo..."):
                en = llamar_ia(modelo_seleccionado, f"Traduce: '{p}'", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="").replace('"', '').replace("'", "").strip()
                url_img = f"https://image.pollinations.ai/prompt/{urllib.parse.quote(en)}?width=1024&height=1024&nologo=true"
                st.image(url_img)
                st.session_state.galeria_imgs.append(url_img)
    elif sub_idx == 2:
        m = st.text_input("Marca:")
        if st.button("Logo") and m:
            url_img = f"https://image.pollinations.ai/prompt/{urllib.parse.quote(f'logo {m}, minimalist, white background')}?width=1024&height=1024&nologo=true"
            st.image(url_img)
            st.session_state.galeria_imgs.append(url_img)
    elif sub_idx == 3:
        u = st.text_input("URL YouTube/TikTok:")
        if st.button("⬇️ Descargar") and u and HAS_YTDLP:
            with st.spinner("Descargando..."):
                t = tempfile.mkdtemp(); y = {'outtmpl': os.path.join(t, "%(title)s.%(ext)s"), 'quiet': True, 'format': 'bestvideo[height<=720]+bestaudio/best'}
                try:
                    with yt_dlp.YoutubeDL(y) as ydl:
                        i = ydl.extract_info(u, download=True); f = ydl.prepare_filename(i)
                        with open(f, "rb") as file: st.session_state.dl_bytes = file.read()
                        st.session_state.dl_filename = os.path.basename(f); st.success("✅ Descargado.")
                except Exception as e: st.error(e)
                finally:
                    for f in os.listdir(t): os.remove(os.path.join(t, f))
                    os.rmdir(t)
        if st.session_state.dl_bytes: st.download_button("💾 Guardar", st.session_state.dl_bytes, st.session_state.dl_filename)
    elif sub_idx == 4:
        st.markdown("#### 🖼️ Galería de Imágenes Generadas")
        if not st.session_state.galeria_imgs:
            st.info("Aún no has generado ninguna imagen en esta sesión.")
        else:
            for idx, url in enumerate(reversed(st.session_state.galeria_imgs)):
                st.image(url, caption=f"Imagen {idx+1}")

# ==========================================
# CATEGORÍA 5: ESTUDIO Y OFICIOS
# ==========================================
elif macro_idx == 5:
    sub_names = ["🌐 REDES", "💇‍♀️ PELUQUERÍA", "📚 DEBERES", "🦺 PRL"]
    sub_idx = custom_tabs(sub_names, "study_sub", 0)
    if sub_idx == 0:
        c1, c2 = st.columns(2)
        with c1:
            ip = st.text_input("IP/CIDR:")
            if st.button("Calcular") and ip: st.info(llamar_ia(modelo_seleccionado, f"Red {ip}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic=""))
        with c2:
            t = campo_con_ideas("Tarea:", "net", "redes", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Ver IP", "", "")
            if st.button("Comando") and t: st.info(llamar_ia(modelo_seleccionado, f"Comando: {t}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic=""))
    elif sub_idx == 1:
        a = st.text_input("Actual:"); d = st.text_input("Deseado:")
        if st.button("Fórmula") and a and d: st.info(llamar_ia(modelo_seleccionado, f"De {a} a {d}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic=""))
    elif sub_idx == 2:
        d = campo_con_ideas("Duda:", "home", "deberes", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Fotosíntesis", "", "")
        if st.button("Explicar") and d: st.info(llamar_ia(modelo_seleccionado, d, st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic=""))
    elif sub_idx == 3:
        t = st.selectbox("Tipo:", ["Charla", "Evaluación", "EPIs", "Legal"]); c = campo_con_ideas("Contexto:", "prl", "riesgos", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Altura", "", "")
        if st.button("Generar") and c:
            res = llamar_ia(modelo_seleccionado, f"PRL {t}: {c}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.text_area("Resultado:", res, height=300); botones_descarga(res, "PRL")

# ==========================================
# CATEGORÍA 6: VIAJES Y HOGAR
# ==========================================
elif macro_idx == 6:
    sub_names = ["🌍 VIAJES", "💰 FINANZAS", "✉️ CARTAS"]
    sub_idx = custom_tabs(sub_names, "travel_sub", 0)
    if sub_idx == 0:
        d = campo_con_ideas("Destino:", "travel", "viajes", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Roma", "", "")
        if st.button("Itinerario") and d:
            res = llamar_ia(modelo_seleccionado, f"Itinerario {d} 3 días", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.info(res); botones_descarga(res, f"Viaje_{d}")
    elif sub_idx == 1:
        i = st.number_input("Ingresos:"); g = st.number_input("Gastos:")
        if st.button("Ahorro"): st.success(f"Ahorras {i-g}€") if i > g else st.error(f"Faltan {g-i}€")
    elif sub_idx == 2:
        t = st.selectbox("Tipo:", ["Reyes", "Discurso", "Reclamación"]); d = campo_con_ideas("Detalles:", "letter", "cartas", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Carta a padres", "", "")
        if st.button("Redactar") and d:
            res = llamar_ia(modelo_seleccionado, f"{t}: {d}", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.text_area("Resultado:", res, height=300); botones_descarga(res, "Carta")

# ==========================================
# CATEGORÍA 7: GESTIÓN DE PLANTAS (NIVEL PRO)
# ==========================================
elif macro_idx == 7:
    if st.session_state.get("rol") == "ninos":
        st.error("🔒 Esta sección está bloqueada en Modo Niños.")
        st.stop()
        
    sub_names = ["📊 DASHBOARD", "⚖️ BÁSCULA", "🔄 TRAZABILIDAD", "🔧 MANTENIMIENTO", "🏭 SIMULADOR", "📄 INFORMES"]
    sub_idx = custom_tabs(sub_names, "plant_sub", 0)
    
    if sub_idx == 0:
        st.markdown("#### 📊 Dashboard de Planta")
        
        # --- ALTA DE VEHÍCULOS ---
        with st.expander("📋 Alta de Vehículos y Conductores (Haz clic para abrir)", expanded=False):
            with st.form("form_alta_vehiculo"):
                c1, c2 = st.columns(2)
                with c1:
                    mat_nueva = st.text_input("Matrícula Camión (Ej: 1234 ABC):").upper()
                    mat_remolque = st.text_input("Matrícula Remolque (Opcional):").upper()
                    empresa_nueva = st.text_input("Empresa de Transporte:")
                with c2:
                    conductor_nuevo = st.text_input("Nombre del Conductor:")
                    dni_nuevo = st.text_input("DNI / Pasaporte:")
                    tara_nueva = st.number_input("Tara Fija del Camión (kg):", min_value=0.0, value=12000.0, step=100.0)
                
                btn_alta = st.form_submit_button("💾 Guardar Vehículo en Flota", type="primary")
                if btn_alta and mat_nueva and empresa_nueva:
                    try:
                        insert_db("INSERT INTO flota_vehiculos (matricula, empresa, conductor, dni, tara_fija) VALUES (?, ?, ?, ?, ?)", 
                                  (mat_nueva, empresa_nueva, conductor_nuevo, dni_nuevo, tara_nueva))
                        st.success(f"✅ Vehículo {mat_nueva} registrado correctamente.")
                        time.sleep(1); st.rerun()
                    except Exception as e:
                        st.error(f"⚠️ Error: Es posible que la matrícula ya esté registrada. Detalle: {e}")
        
        st.markdown("---")
        df_bascula = get_db_data("bascula"); df_lotes = get_db_data("lotes"); df_mant = get_db_data("mantenimiento"); df_flota = get_db_data("flota_vehiculos")
        
        c1, c2, c3, c4 = st.columns(4)
        with c1: st.metric("Vehículos en Flota", len(df_flota))
        with c2: 
            if not df_bascula.empty: st.metric("Total Entradas (kg)", f"{df_bascula[df_bascula['operacion'] == 'Entrada']['peso_neto'].sum():,.0f}")
            else: st.metric("Total Entradas", "0 kg")
        with c3:
            if not df_lotes.empty: st.metric("Lotes Activos", len(df_lotes[df_lotes['estado'] != 'Completado']))
            else: st.metric("Lotes Activos", "0")
        with c4:
            if not df_mant.empty: st.metric("Máquinas Averiadas", len(df_mant[df_mant['estado'] == 'Averiada']), delta_color="inverse")
            else: st.metric("Máquinas Averiadas", "0")
            
        st.markdown("---")
        
        # --- TABLA DE FLOTA ---
        st.markdown("##### Flota de Vehículos Registrados")
        if not df_flota.empty:
            st.dataframe(df_flota[['matricula', 'empresa', 'conductor', 'dni', 'tara_fija']], use_container_width=True, hide_index=True)
        else:
            st.info("No hay vehículos en la flota. Registra uno arriba.")

    elif sub_idx == 1:
        st.markdown("#### ⚖️ Control de Báscula (Nivel Avanzado)")
        
        df_flota = get_db_data("flota_vehiculos")
        if df_flota.empty:
            st.warning("⚠️ No hay vehículos registrados. Ve al Dashboard y registra un vehículo primero.")
        else:
            # Selector de matrícula FUERA del form para que actualice los datos en tiempo real
            lista_matriculas = df_flota['matricula'].tolist()
            mat_sel = st.selectbox("1. Selecciona el Camión (Matrícula):", lista_matriculas)
            
            # Autocompletar datos del vehículo
            row_veh = df_flota[df_flota['matricula'] == mat_sel].iloc[0]
            st.info(f"**Empresa:** {row_veh['empresa']} | **Conductor:** {row_veh['conductor']} | **DNI:** {row_veh['dni']} | **Tara Base:** {row_veh['tara_fija']} kg")
            
            with st.form("form_bascula_pro"):
                st.markdown("##### 2. Datos de la Carga")
                c1, c2 = st.columns(2)
                with c1:
                    residuo = st.selectbox("Tipo de Residuo:", list(CATALOGO_RESIDUOS.keys()))
                    ler_sel = st.selectbox("Código LER:", list(CATALOGO_LER.keys()), format_func=lambda x: f"{x} - {CATALOGO_LER[x]}")
                    origen = st.text_input("Origen de la Mercancía:", placeholder="Ej: Ayuntamiento de Valencia")
                with c2:
                    destino = st.selectbox("Destino Interno:", DESTINOS_INTERNOS)
                    albaran = st.text_input("Nº Albarán del Cliente:")
                    doc_di = st.text_input("Nº Documento DI:")
                
                st.markdown("##### 3. Control de Calidad y Documentación")
                inspeccion_ok = st.checkbox("Carga Apta (Sin impropios ni exceso de humedad)", value=True)
                incidencia = ""
                penalizacion = 0.0
                if not inspeccion_ok:
                    incidencia = st.text_area("Describe la incidencia:")
                    penalizacion = st.number_input("Penalización / Merma (%):", min_value=0.0, max_value=100.0, value=10.0, step=1.0)
                
                st.markdown("##### 4. Pesaje")
                c3, c4, c5 = st.columns(3)
                with c3:
                    operacion = st.radio("Operación:", ["Entrada", "Salida"], horizontal=True)
                with c4:
                    peso_bruto = st.number_input("Peso Bruto (kg):", min_value=0.0, value=0.0, step=10.0)
                with c5:
                    tara_real = st.number_input("Tara (kg):", min_value=0.0, value=float(row_veh['tara_fija']), step=10.0)
                
                btn_save = st.form_submit_button("💾 Guardar Pesaje en Báscula", type="primary")
                
                if btn_save:
                    if peso_bruto <= 0:
                        st.error("⚠️ Debes introducir un Peso Bruto válido.")
                    else:
                        # Cálculos
                        peso_neto_bruto = peso_bruto - tara_real
                        merma_kg = peso_neto_bruto * (penalizacion / 100)
                        peso_neto_final = peso_neto_bruto - merma_kg
                        
                        fecha_actual = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        insert_db("""INSERT INTO bascula 
                                     (fecha, matricula, empresa, residuo, codigo_ler, origen, destino, peso_bruto, tara, peso_neto, penalizacion, albaran, doc_di, inspeccion_ok, incidencia, operacion) 
                                     VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", 
                                  (fecha_actual, mat_sel, row_veh['empresa'], residuo, ler_sel, origen, destino, peso_bruto, tara_real, peso_neto_final, penalizacion, albaran, doc_di, int(inspeccion_ok), incidencia, operacion))
                        
                        st.success(f"✅ {operacion} registrada. Peso Neto Facturable: {peso_neto_final} kg")
                        time.sleep(1); st.rerun()

            # --- HISTORIAL DE PESAJE ---
            st.markdown("---")
            st.markdown("##### Últimos Pesajes Registrados")
            df_hist = get_db_data("bascula").sort_values("id", ascending=False).head(10)
            if not df_hist.empty: 
                # Mostramos las columnas más relevantes
                cols_mostrar = ['fecha', 'matricula', 'empresa', 'operacion', 'residuo', 'codigo_ler', 'peso_neto', 'albaran']
                st.dataframe(df_hist[cols_mostrar], use_container_width=True)
            else: 
                st.info("No hay registros en la báscula todavía.")

    elif sub_idx == 2:
        st.markdown("#### 🔄 Trazabilidad de Lotes")
        
        # Lógica para sacar los camiones que están DENTRO (Han entrado pero no han salido)
        df_bascula = get_db_data("bascula")
        camiones_dentro = []
        if not df_bascula.empty:
            # Filtramos solo los que tienen una entrada reciente sin salida posterior
            for mat in df_bascula['matricula'].unique():
                df_mat = df_bascula[df_bascula['matricula'] == mat].sort_values('id', ascending=False)
                if not df_mat.empty and df_mat.iloc[0]['operacion'] == 'Entrada':
                    camiones_dentro.append(mat)

        with st.form("form_lote"):
            c1, c2 = st.columns(2)
            with c1:
                if camiones_dentro:
                    lote_matricula = st.selectbox("Selecciona el Camión en Planta (Matrícula):", ["Selecciona una..."] + camiones_dentro)
                else:
                    st.warning("No hay camiones dentro de la planta actualmente.")
                    lote_matricula = None
                
                lote_residuo = st.selectbox("Residuo del Lote:", list(CATALOGO_RESIDUOS.keys()))
                lote_peso_in = st.number_input("Peso Inicial del Lote (kg):", min_value=0.0, value=5000.0, step=100.0)
            with c2:
                lote_estado = st.selectbox("Estado Inicial:", ["Recibido", "En Proceso", "Clasificado"])
                btn_lote = st.form_submit_button("Crear Lote", type="primary")
                
                if btn_lote:
                    if not lote_matricula or lote_matricula == "Selecciona una...":
                        st.error("⚠️ Debes seleccionar una matrícula válida.")
                    else:
                        fecha = datetime.datetime.now().strftime("%Y-%m-%d")
                        insert_db("INSERT INTO lotes (fecha, residuo, peso_in, estado, peso_out) VALUES (?, ?, ?, ?, ?)", 
                                  (fecha, f"{lote_residuo} - Camión: {lote_matricula}", lote_peso_in, lote_estado, 0.0))
                        st.success("✅ Lote creado y vinculado al camión."); time.sleep(1); st.rerun()

        st.markdown("---")
        st.markdown("##### Lotes Activos")
        df_lotes = get_db_data("lotes")
        if not df_lotes.empty:
            for idx, row in df_lotes.iterrows():
                if row['estado'] != 'Completado':
                    with st.expander(f"Lote #{row['id']} - {row['residuo']} ({row['fecha']})"):
                        st.write(f"Peso Inicial: {row['peso_in']} kg")
                        nuevo_estado = st.selectbox("Actualizar Estado:", ["Recibido", "En Proceso", "Clasificado", "Completado"], index=0, key=f"lote_est_{row['id']}")
                        peso_out = st.number_input("Peso Final Recuperado (kg):", value=row['peso_out'], key=f"lote_p_{row['id']}")
                        if st.button("Actualizar Lote", key=f"btn_lote_{row['id']}"):
                            insert_db("UPDATE lotes SET estado=?, peso_out=? WHERE id=?", (nuevo_estado, peso_out, row['id']))
                            st.success("Actualizado."); time.sleep(1); st.rerun()
        else: st.info("No hay lotes registrados.")

    elif sub_idx == 3:
        st.markdown("#### 🔧 Mantenimiento y Averías")
        with st.form("form_mant"):
            c1, c2 = st.columns(2)
            with c1:
                maquina = st.selectbox("Máquina:", ["Trituradora Primaria", "Cinta Transportadora 1", "Separador Magnético", "Separador Balístico", "Prensa Hidráulica", "Otra"])
                descripcion = st.text_area("Descripción del fallo:")
            with c2:
                st.markdown("<br><br><br><br>", unsafe_allow_html=True)
                btn_averia = st.form_submit_button("🚨 Reportar Avería", type="primary")
                if btn_averia and descripcion:
                    fecha = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
                    insert_db("INSERT INTO mantenimiento (fecha, maquina, descripcion, estado) VALUES (?, ?, ?, ?)", (fecha, maquina, descripcion, "Averiada"))
                    st.error("Avería reportada. La planta está parada en ese sector."); time.sleep(1); st.rerun()
        st.markdown("---")
        st.markdown("##### Estado de las Máquinas")
        df_mant = get_db_data("mantenimiento")
        if not df_mant.empty:
            for idx, row in df_mant.iterrows():
                if row['estado'] == 'Averiada':
                    with st.expander(f"🚨 {row['maquina']} - {row['fecha']}", expanded=True):
                        st.write(f"Fallo: {row['descripcion']}")
                        if st.button("✅ Marcar como Reparada", key=f"rep_{row['id']}"):
                            insert_db("UPDATE mantenimiento SET estado='Reparada' WHERE id=?", (row['id'],))
                            st.success("Máquina reparada."); time.sleep(1); st.rerun()
                else: st.success(f"✅ {row['maquina']} - Funcionando (Reparada el {row['fecha']})")
        else: st.info("Todas las máquinas funcionan correctamente.")

    elif sub_idx == 4:
        sim_sub_names = ["🧮 SIMULAR", "📦 INVENTARIO", "🔧 MANTENIMIENTO LEGAL", "📊 COMPARAR"]
        sim_sub_idx = custom_tabs(sim_sub_names, "plant_sim_sub", 0)
        
        if sim_sub_idx == 0:
            render_simular()
        elif sim_sub_idx == 1:
            render_inventario()
        elif sim_sub_idx == 2:
            render_mantenimiento_legal()
        elif sim_sub_idx == 3:
            render_comparar()

    elif sub_idx == 5:
        render_informe()

# ==========================================
# CATEGORÍA 8: MÚSICA Y ARTE
# ==========================================
elif macro_idx == 8:
    sub_names = ["🎵 CANCIONES", "🎨 POESÍA"]
    sub_idx = custom_tabs(sub_names, "music_sub", 0)
    if sub_idx == 0:
        t = campo_con_ideas("Tema:", "song", "canciones", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Amor", "", "")
        g = st.selectbox("Género:", ["Pop", "Rock", "Rap"]); to = st.select_slider("Tono:", ["Triste", "Alegre", "Épico"])
        if st.button("Escribir") and t:
            letra = llamar_ia(modelo_seleccionado, f"Canción {g} sobre {t}, tono {to}.", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in letra: st.error(letra)
            else:
                st.text_area("Letra:", letra, height=300); botones_descarga(letra, "Cancion")
                if st.button("Audio"):
                    aud = generar_audio_sincrono(letra, "es-ES-ElviraNeural")
                    if aud: st.audio(aud, format="audio/mp3")
    elif sub_idx == 1:
        t = campo_con_ideas("Tema:", "poem", "poemas", modelo_seleccionado, st.session_state.api_key_gemini_manual, "Ej: Mar", "", "")
        if st.button("Escribir") and t:
            res = llamar_ia(modelo_seleccionado, f"Poema sobre {t}.", st.session_state.api_key_gemini_manual, api_key_openai="", api_key_anthropic="")
            if "[ERROR" in res: st.error(res)
            else: st.text_area("Poema:", res, height=300); botones_descarga(res, "Poema")
